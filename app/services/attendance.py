"""Attendance storage is deliberately pinned to the Slurry primary database."""
from datetime import date, datetime, timedelta
from io import BytesIO
from uuid import uuid4
import bcrypt

from ..db import transaction
from ..departments import department_context


def _run(callback):
    with department_context('slurry'):
        with transaction() as cursor:
            return callback(cursor)


def _normalise_employee_row(row):
    """Convert MySQL integer flags to JSON-safe Boolean values."""
    row['pin_reset_required'] = bool(row.get('pin_reset_required'))
    if 'is_active' in row:
        row['is_active'] = bool(row.get('is_active'))
    return row


def authenticate_employee(pin):
    """Return the active employee whose bcrypt hash matches a four-digit PIN."""
    pin = str(pin or '')
    if not pin.isdigit() or len(pin) != 4:
        raise ValueError('PIN must contain exactly 4 digits.')

    def query(cursor):
        cursor.execute(
            "SELECT id,employee_code,employee_name,staffing_company,department,"
            "job_title,pin_hash,pin_reset_required "
            "FROM system_import.attendance_employees WHERE is_active=TRUE"
        )
        return cursor.fetchall()

    rows = _run(query)
    row = next(
        (
            candidate
            for candidate in rows
            if bcrypt.checkpw(pin.encode(), candidate['pin_hash'].encode())
        ),
        None,
    )
    if not row:
        raise ValueError('Invalid PIN.')
    row.pop('pin_hash')
    return _normalise_employee_row(row)

def employee_history(employee_id):
    def query(cursor):
        cursor.execute(
            "SELECT id,attendance_date,paid_break_hours,total_hours,"
            "overtime_hours,status,employee_note,manager_note "
            "FROM system_import.attendance_records WHERE employee_id=%s "
            "ORDER BY attendance_date DESC LIMIT 14",
            (employee_id,),
        )
        return _serialise_attendance_rows(cursor.fetchall())

    return _run(query)

def submit_attendance(employee_id, payload):
    paid_break = _half_hour(
        payload.get('paid_break_hours'), 'Paid break hours', 0, 2
    )
    total_hours = _half_hour(payload.get('total_hours'), 'Total hours', 3, 12)
    overtime = _half_hour(
        payload.get('overtime_hours'),
        'Overtime hours',
        0.5,
        24,
        allow_blank=True,
    )
    today = datetime.now().date()

    def save(cursor):
        cursor.execute(
            "SELECT id FROM system_import.attendance_records "
            "WHERE employee_id=%s AND attendance_date=%s FOR UPDATE",
            (employee_id, today),
        )
        row = cursor.fetchone()
        if row:
            cursor.execute(
                "UPDATE system_import.attendance_records "
                "SET paid_break_hours=%s,total_hours=%s,overtime_hours=%s,"
                "status='present',row_version=row_version+1 WHERE id=%s",
                (paid_break, total_hours, overtime, row['id']),
            )
        else:
            cursor.execute(
                "INSERT INTO system_import.attendance_records("
                "employee_id,attendance_date,paid_break_hours,total_hours,"
                "overtime_hours,status) VALUES(%s,%s,%s,%s,%s,'present')",
                (employee_id, today, paid_break, total_hours, overtime),
            )

    _run(save)
    return today.isoformat()

def list_employees(page=1, page_size=50, query=None):
    term = str(query or '').strip()
    where, params = '', []
    if term:
        where = " WHERE staffing_company LIKE %s OR employee_name LIKE %s"
        params = [f"%{term}%", f"%{term}%"]
    offset = (page - 1) * page_size

    def fetch(cursor):
        cursor.execute(
            "SELECT COUNT(*) total FROM system_import.attendance_employees" + where,
            params,
        )
        total = cursor.fetchone()['total']
        cursor.execute(
            "SELECT id,employee_name,staffing_company,department,"
            "job_title,pin_code,pin_reset_required,is_active,created_at,updated_at "
            "FROM system_import.attendance_employees"
            + where
            + " ORDER BY staffing_company,employee_name LIMIT %s OFFSET %s",
            [*params, page_size, offset],
        )
        employees = [
            _normalise_employee_row(row)
            for row in cursor.fetchall()
        ]
        return employees, total

    return _run(fetch)

def save_employee(payload, employee_id=None):
    """Create or edit an employee; a manager-supplied PIN is always temporary."""
    company = str(payload.get('staffing_company', '')).strip()
    name = str(payload.get('employee_name', '')).strip()
    pin = str(payload.get('pin', ''))
    if not company or not name:
        raise ValueError('Staffing company and name are required.')
    if pin and (not pin.isdigit() or len(pin) != 4):
        raise ValueError('PIN must contain exactly 4 digits.')
    if not employee_id and not pin:
        raise ValueError('A PIN is required for a new employee.')

    pin_hash = bcrypt.hashpw(pin.encode(), bcrypt.gensalt()).decode() if pin else None
    active = bool(payload.get('is_active', True))
    department = str(
        payload.get('employee_department', payload.get('department', ''))
    ).strip() or None
    title = str(payload.get('job_title', '')).strip() or None

    def save(cursor):
        if pin:
            cursor.execute(
                "SELECT id,pin_hash FROM system_import.attendance_employees "
                "WHERE id<>%s FOR UPDATE",
                (employee_id or 0,),
            )
            if any(
                bcrypt.checkpw(pin.encode(), row['pin_hash'].encode())
                for row in cursor.fetchall()
            ):
                raise ValueError(
                    'PIN is already assigned to another employee. Choose a unique PIN.'
                )
        if employee_id:
            # Leaving PIN blank preserves it. Supplying one forces first-login setup.
            cursor.execute(
                "UPDATE system_import.attendance_employees SET employee_name=%s,"
                "staffing_company=%s,department=%s,job_title=%s,is_active=%s,"
                "pin_hash=COALESCE(%s,pin_hash),pin_code=COALESCE(%s,pin_code),"
                "pin_reset_required=IF(%s IS NULL,pin_reset_required,TRUE) "
                "WHERE id=%s",
                (
                    name,
                    company,
                    department,
                    title,
                    active,
                    pin_hash,
                    pin or None,
                    pin_hash,
                    employee_id,
                ),
            )
            if cursor.rowcount != 1:
                raise ValueError('Employee not found.')
        else:
            code = 'ATT-' + uuid4().hex.upper()
            cursor.execute(
                "INSERT INTO system_import.attendance_employees("
                "employee_code,employee_name,staffing_company,department,"
                "job_title,pin_hash,pin_code,pin_reset_required,is_active) "
                "VALUES(%s,%s,%s,%s,%s,%s,%s,TRUE,%s)",
                (code, name, company, department, title, pin_hash, pin, active),
            )

    try:
        _run(save)
    except Exception as exc:
        if 'Duplicate' in str(exc):
            raise ValueError(
                'Could not create a unique internal employee record. Try again.'
            ) from exc
        raise

def reset_employee_pin(employee_id, new_pin, confirm_pin):
    """Replace a temporary manager PIN exactly once with an employee PIN."""
    new_pin = str(new_pin or '')
    confirm_pin = str(confirm_pin or '')
    if not new_pin.isdigit() or len(new_pin) != 4:
        raise ValueError('New PIN must contain exactly 4 digits.')
    if new_pin != confirm_pin:
        raise ValueError('The new PIN and confirmation do not match.')
    new_hash = bcrypt.hashpw(new_pin.encode(), bcrypt.gensalt()).decode()

    def save(cursor):
        cursor.execute(
            "SELECT id,pin_hash,pin_reset_required "
            "FROM system_import.attendance_employees "
            "WHERE id=%s AND is_active=TRUE FOR UPDATE",
            (employee_id,),
        )
        employee = cursor.fetchone()
        if not employee:
            raise ValueError('Employee not found or inactive.')
        if not employee['pin_reset_required']:
            raise ValueError('First-login PIN setup has already been completed.')
        if bcrypt.checkpw(new_pin.encode(), employee['pin_hash'].encode()):
            raise ValueError(
                'Choose a new PIN that is different from the temporary PIN.'
            )
        cursor.execute(
            "SELECT id,pin_hash FROM system_import.attendance_employees WHERE id<>%s FOR UPDATE",
            (employee_id,),
        )
        if any(
            bcrypt.checkpw(new_pin.encode(), row['pin_hash'].encode())
            for row in cursor.fetchall()
        ):
            raise ValueError(
                'PIN is already assigned to another employee. Choose a unique PIN.'
            )
        cursor.execute(
            "UPDATE system_import.attendance_employees "
            "SET pin_hash=%s,pin_code=%s,pin_reset_required=FALSE WHERE id=%s",
            (new_hash, new_pin, employee_id),
        )

    _run(save)

def delete_employee(employee_id):
    """Remove unused employees; archive employees that already own attendance history."""

    def remove(cursor):
        cursor.execute(
            "SELECT COUNT(*) total FROM system_import.attendance_records "
            "WHERE employee_id=%s",
            (employee_id,),
        )
        has_history = cursor.fetchone()['total'] > 0
        if has_history:
            cursor.execute(
                "UPDATE system_import.attendance_employees "
                "SET is_active=FALSE WHERE id=%s",
                (employee_id,),
            )
            action = 'archived'
        else:
            cursor.execute(
                "DELETE FROM system_import.attendance_employees WHERE id=%s",
                (employee_id,),
            )
            action = 'deleted'
        if cursor.rowcount != 1:
            raise ValueError('Employee not found.')
        return action

    return _run(remove)


def _filters(
    staffing_company=None,
    employee_name=None,
    start_date=None,
    end_date=None,
):
    clauses = []
    params = []
    if staffing_company:
        clauses.append('e.staffing_company LIKE %s')
        params.append(f'%{staffing_company}%')
    if employee_name:
        clauses.append('e.employee_name LIKE %s')
        params.append(f'%{employee_name}%')
    date_filters = [
        ('r.attendance_date >= %s', start_date),
        ('r.attendance_date <= %s', end_date),
    ]
    for column, value in date_filters:
        if value:
            try:
                date.fromisoformat(value)
            except ValueError as exc:
                raise ValueError('Choose valid attendance dates.') from exc
            clauses.append(column)
            params.append(value)
    return (' WHERE ' + ' AND '.join(clauses) if clauses else ''), params


def _serialise_attendance_rows(rows):
    for row in rows:
        if isinstance(row.get('attendance_date'), date):
            row['attendance_date'] = row['attendance_date'].isoformat()
        for key in ('clock_in', 'clock_out'):
            if isinstance(row.get(key), datetime):
                row[key] = row[key].isoformat(timespec='minutes')
    return rows

def _half_hour(raw, label, minimum, maximum, allow_blank=False):
    if raw in (None, ''):
        if allow_blank:
            return None
        raise ValueError(f'Choose {label.lower()}.')
    try:
        value = float(raw)
    except (TypeError, ValueError) as exc:
        raise ValueError(f'{label} must be a number.') from exc
    invalid_increment = abs(value * 2 - round(value * 2)) > 1e-9
    if value < minimum or value > maximum or invalid_increment:
        raise ValueError(
            f'{label} must be between {minimum:.1f} and {maximum:.1f} '
            'in 0.5-hour increments.'
        )
    return value


def list_attendance(page=1, page_size=50, **filters):
    where, params = _filters(**filters)
    offset = (page - 1) * page_size

    def query(cursor):
        cursor.execute(
            "SELECT COUNT(*) total FROM system_import.attendance_records r "
            "JOIN system_import.attendance_employees e ON e.id=r.employee_id"
            + where,
            params,
        )
        total = cursor.fetchone()['total']
        cursor.execute(
            "SELECT r.id,r.employee_id,e.employee_name,e.staffing_company,"
            "e.department,e.job_title,r.attendance_date,r.clock_in,r.clock_out,"
            "r.paid_break_hours,r.total_hours,r.overtime_hours,r.status,"
            "r.employee_note,r.manager_note,r.row_version "
            "FROM system_import.attendance_records r "
            "JOIN system_import.attendance_employees e ON e.id=r.employee_id"
            + where
            + " ORDER BY r.attendance_date DESC,e.staffing_company,"
            "e.employee_name LIMIT %s OFFSET %s",
            [*params, page_size, offset],
        )
        return _serialise_attendance_rows(cursor.fetchall()), total

    return _run(query)

def update_attendance(record_id, payload):
    status = str(payload.get('status', '')).lower()
    if status not in {'present', 'absent', 'leave', 'excused'}:
        raise ValueError('Choose a valid attendance status.')
    note = str(payload.get('manager_note', '')).strip()
    clock_in = payload.get('clock_in') or None
    clock_out = payload.get('clock_out') or None
    paid_break = _half_hour(
        payload.get('paid_break_hours'),
        'Paid break hours',
        0,
        2,
        allow_blank=True,
    )
    total_hours = _half_hour(
        payload.get('total_hours'),
        'Total hours',
        3,
        12,
        allow_blank=True,
    )
    overtime = _half_hour(
        payload.get('overtime_hours'),
        'Overtime hours',
        0.5,
        24,
        allow_blank=True,
    )

    def save(cursor):
        cursor.execute(
            "UPDATE system_import.attendance_records SET clock_in=%s,clock_out=%s,"
            "paid_break_hours=%s,total_hours=%s,overtime_hours=%s,status=%s,"
            "manager_note=%s,row_version=row_version+1 "
            "WHERE id=%s AND row_version=%s",
            (
                clock_in,
                clock_out,
                paid_break,
                total_hours,
                overtime,
                status,
                note or None,
                record_id,
                int(payload.get('row_version', 0)),
            ),
        )
        if cursor.rowcount != 1:
            raise ValueError(
                'Record changed or was not found. Refresh and try again.'
            )

    _run(save)

def export_attendance(**filters):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError as exc:
        raise ValueError(
            'Excel export is unavailable. Run pip install -r requirements.txt '
            'in the Python 3.13 environment.'
        ) from exc

    rows, _ = list_attendance(page=1, page_size=100000, **filters)
    record_dates = sorted(
        {
            date.fromisoformat(row['attendance_date'])
            for row in rows
            if row.get('attendance_date')
        }
    )
    start_date = filters.get('start_date')
    end_date = filters.get('end_date')
    if start_date and end_date:
        start = date.fromisoformat(start_date)
        end = date.fromisoformat(end_date)
        record_dates = [
            start + timedelta(days=offset)
            for offset in range((end - start).days + 1)
        ]

    # Pivot records into one employee row and one calendar-date column.
    employees = {}
    for row in rows:
        employee_key = (row['staffing_company'], row['employee_name'])
        employee = employees.setdefault(
            employee_key,
            {
                'staffing_company': row['staffing_company'],
                'employee_name': row['employee_name'],
                'hours': {},
            },
        )
        worked_hours = (
            float(row.get('total_hours') or 0)
            - float(row.get('paid_break_hours') or 0)
            + float(row.get('overtime_hours') or 0)
        )
        employee['hours'][date.fromisoformat(row['attendance_date'])] = round(
            worked_hours,
            2,
        )

    book = Workbook()
    sheet = book.active
    sheet.title = 'Attendance'
    sheet.freeze_panes = 'C3'

    sheet.merge_cells('A1:A2')
    sheet.merge_cells('B1:B2')
    sheet['A1'] = 'Staffing company / 劳务公司'
    sheet['B1'] = 'Employee name / 员工姓名'
    for column, attendance_date in enumerate(record_dates, start=3):
        date_cell = sheet.cell(row=1, column=column, value=attendance_date)
        date_cell.number_format = 'd'
        weekday = attendance_date.strftime('%a')
        if attendance_date.weekday() >= 5:
            weekday += ' / Weekend'
        sheet.cell(row=2, column=column, value=weekday)

    for employee in employees.values():
        sheet.append([
            employee['staffing_company'],
            employee['employee_name'],
            *[employee['hours'].get(day) for day in record_dates],
        ])

    header_fill = PatternFill('solid', fgColor='0070C0')
    weekend_fill = PatternFill('solid', fgColor='5B9BD5')
    header_font = Font(color='FFFFFF', bold=True)
    for row_number in (1, 2):
        for cell in sheet[row_number]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
    for column, attendance_date in enumerate(record_dates, start=3):
        if attendance_date.weekday() >= 5:
            sheet.cell(row=1, column=column).fill = weekend_fill
            sheet.cell(row=2, column=column).fill = weekend_fill

    sheet.column_dimensions['A'].width = 24
    sheet.column_dimensions['B'].width = 24
    for column in range(3, len(record_dates) + 3):
        sheet.column_dimensions[sheet.cell(row=1, column=column).column_letter].width = 15

    output = BytesIO()
    book.save(output)
    output.seek(0)
    return output
