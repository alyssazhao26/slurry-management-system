import hashlib
import io
import json
import re
from datetime import date, datetime, timedelta
from uuid import UUID, uuid4

from flask import current_app

from ..db import transaction
from ..departments import current_department, uses_slitting_workflow
from . import slitting
from .analysis import Analysis, analyse_abnormality, analyse_production


def _json_safe_rows(rows):
    """Convert MySQL TIME values to HH:MM before Flask returns JSON."""
    for row in rows:
        for key, value in row.items():
            if isinstance(value, timedelta):
                total_minutes = int(value.total_seconds() // 60)
                hours, minutes = divmod(total_minutes, 60)
                row[key] = f"{hours:02d}:{minutes:02d}"
            elif isinstance(value, datetime):
                row[key] = value.isoformat(sep=" ", timespec="seconds")
            elif isinstance(value, date):
                row[key] = value.isoformat()
    return rows


def _downtime_minutes(start_time, end_time):
    """Calculate elapsed minutes, treating an earlier end time as the next day."""
    time_pattern = re.compile(r"^([01]\d|2[0-3]):[0-5]\d$")
    start = str(start_time or "")[:5]
    end = str(end_time or "")[:5]
    if not time_pattern.fullmatch(start) or not time_pattern.fullmatch(end):
        raise ValueError("Event time must use valid 24-hour values (HH:MM - HH:MM).")
    start_hour, start_minute = (int(part) for part in start.split(":"))
    end_hour, end_minute = (int(part) for part in end.split(":"))
    minutes = end_hour * 60 + end_minute - (start_hour * 60 + start_minute)
    return minutes if minutes >= 0 else minutes + 24 * 60


def _audit(cursor, actor_id, entity_type, entity_id, action, details):
    cursor.execute(
        """
        INSERT INTO audit_events (actor_id, entity_type, entity_id, action, details_json)
        VALUES (%s, %s, %s, %s, %s)
        """,
        (actor_id, entity_type, entity_id, action, json.dumps(details)),
    )


def _validate_active_event_type(cursor, payload):
    if uses_slitting_workflow():
        cursor.execute(
            """SELECT 1 FROM system_import.slitting_event_categories
            WHERE main_category=%s AND subcategory=%s AND is_active=TRUE LIMIT 1""",
            (payload.get("event_type"), payload.get("subcategory")),
        )
        if not cursor.fetchone():
            raise ValueError("An invalid or inactive Slitting main/subcategory was supplied.")
        return
    cursor.execute(
        "SELECT event_value FROM system_import.event_type_options WHERE is_active = TRUE"
    )
    if payload.get("event_type") not in {row["event_value"] for row in cursor.fetchall()}:
        raise ValueError("An invalid or inactive event type was supplied.")


def _sync_identity(payload, source_type):
    client_record_id = str(payload.get("client_record_id") or uuid4())
    try:
        UUID(client_record_id)
    except ValueError as exc:
        raise ValueError("client_record_id must be a UUID.") from exc

    normalized = dict(payload)
    normalized["client_record_id"] = client_record_id
    payload_json = json.dumps(normalized, sort_keys=True, separators=(",", ":"))
    return client_record_id, hashlib.sha256(payload_json.encode()).hexdigest()


def _existing_sync(cursor, source_type, client_record_id, payload_hash):
    cursor.execute(
        """
        SELECT server_record_id, payload_hash
        FROM system_import.sync_receipts
        WHERE source_type = %s AND client_record_id = %s
        """,
        (source_type, client_record_id),
    )
    receipt = cursor.fetchone()
    if receipt and receipt["payload_hash"] != payload_hash:
        raise ValueError("The same offline record ID was submitted with different data.")
    return receipt["server_record_id"] if receipt else None


def create_production(payload, actor_id=None):
    if uses_slitting_workflow():
        payload = slitting.normalise_production(payload)
    payload = {**payload, "record_date": date.today().isoformat()}
    planned, actual = (float(payload[key]) for key in ("planned_quantity", "actual_quantity"))
    raw_qualified = str(payload.get("qualified_quantity", "")).strip()
    qualified = float(raw_qualified) if raw_qualified else None
    qualified_pending = qualified is None
    if qualified is not None and qualified > actual or actual < 0 or planned < 0:
        raise ValueError(
            "Quantities are invalid: planned and actual must be non-negative, "
            "qualified cannot exceed actual."
        )
    analysis = (
        analyse_production(planned, actual, qualified, current_app.config["LOW_YIELD_THRESHOLD"])
        if qualified is not None
        else Analysis(
            "normal",
            "normal",
            "Qualified quantity is pending follow-up.",
            {"qualified_pending": True},
        )
    )
    achievement_rate = actual / planned if planned > 0 else None
    qualified_rate = qualified / actual if qualified is not None and actual > 0 else None

    formula_code = str(payload["formula_code"]).strip().upper()
    if not formula_code:
        raise ValueError("Choose a formula or enter a custom formula.")
    payload = {**payload, "formula_code": formula_code}
    client_record_id, payload_hash = _sync_identity(payload, "production")
    with transaction() as cursor:
        replay_id = _existing_sync(cursor, "production", client_record_id, payload_hash)
        if replay_id:
            return replay_id, analysis
        cursor.execute(
            """
            SELECT id, qualified_pending, actual_quantity
            FROM system_import.production_records
            WHERE record_date = %s AND shift_name = %s AND machine_code = %s AND batch_number = %s
            """,
            (
                payload["record_date"],
                payload["shift_name"],
                payload["machine_code"],
                payload["batch_number"],
            ),
        )
        existing = cursor.fetchone()
        if existing:
            if uses_slitting_workflow():
                raise ValueError(
                    "This machine, shift and run ID already exist. Use the qualified-roll follow-up or manager edit."
                )
            if existing["qualified_pending"] and qualified is not None:
                existing_actual = float(existing["actual_quantity"])
                if qualified > existing_actual:
                    raise ValueError("Qualified quantity cannot exceed the saved actual quantity.")
                cursor.execute(
                    """
                    UPDATE system_import.production_records
                    SET qualified_quantity = %s, qualified_pending = FALSE,
                        qualified_rate = %s, row_version = row_version + 1
                    WHERE id = %s
                    """,
                    (
                        qualified,
                        qualified / existing_actual if existing_actual > 0 else None,
                        existing["id"],
                    ),
                )
                _audit(
                    cursor,
                    actor_id,
                    "production_record",
                    existing["id"],
                    "qualified_quantity_follow_up",
                    {"qualified_quantity": qualified},
                )
                return existing["id"], analysis
            raise ValueError(
                "This date, shift, machine, and batch already have a production record. "
                "Open Production records to add the qualified result if it is pending."
            )

        cursor.execute(
            """
            INSERT INTO system_import.production_records (
                record_date, shift_name, machine_code, formula_code, batch_number,
                planned_quantity, actual_quantity, qualified_quantity, qualified_pending, notes,
                achievement_rate, qualified_rate, created_by, custom_fields
            ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
            """,
            (
                payload["record_date"],
                payload["shift_name"],
                payload["machine_code"],
                formula_code,
                payload["batch_number"],
                planned,
                actual,
                qualified,
                qualified_pending,
                payload.get("notes", ""),
                achievement_rate,
                qualified_rate,
                actor_id,
                json.dumps(payload.get("custom_fields", {})),
            ),
        )
        record_id = cursor.lastrowid
        cursor.execute(
            """
            INSERT INTO system_import.sync_receipts
                (source_type, client_record_id, payload_hash, server_record_id)
            VALUES ('production', %s, %s, %s)
            """,
            (client_record_id, payload_hash, record_id),
        )
        _audit(cursor, actor_id, "production_record", record_id, "created_kiosk", payload)
        if analysis.status == "open":
            cursor.execute(
                """
                INSERT INTO exceptions_queue
                    (source_type, source_id, severity, summary, evidence_json)
                VALUES ('production', %s, %s, %s, %s)
                """,
                (record_id, analysis.severity, analysis.summary, json.dumps(analysis.evidence)),
            )
    return record_id, analysis


def update_production_qualification(record_id, payload, actor_id=None):
    if uses_slitting_workflow():
        existing = manager_get_record("production", record_id)
        # The operator enters output rolls; the stored qualified quantity is area.
        converted = slitting.normalise_production(
            {"qualified_rolls": payload.get("qualified_quantity")}, existing
        )
        with transaction() as cursor:
            qualified = converted["qualified_quantity"]
            if qualified == "":
                raise ValueError("Enter qualified output rolls.")
            actual = float(existing["actual_quantity"])
            cursor.execute(
                "UPDATE system_import.production_records SET qualified_quantity=%s, qualified_pending=FALSE, qualified_rate=%s, custom_fields=%s, row_version=row_version+1 WHERE id=%s",
                (
                    qualified,
                    qualified / actual if actual else None,
                    json.dumps(converted["custom_fields"]),
                    record_id,
                ),
            )
            _audit(
                cursor,
                actor_id,
                "production_record",
                record_id,
                "qualified_rolls_follow_up",
                converted,
            )
        return record_id
    try:
        qualified = float(payload["qualified_quantity"])
    except (KeyError, TypeError, ValueError) as exc:
        raise ValueError("Enter a valid qualified quantity.") from exc
    if qualified < 0:
        raise ValueError("Qualified quantity cannot be negative.")
    with transaction() as cursor:
        cursor.execute(
            "SELECT planned_quantity, actual_quantity FROM system_import.production_records WHERE id = %s",
            (record_id,),
        )
        row = cursor.fetchone()
        if not row:
            raise ValueError("Production record not found.")
        actual = float(row["actual_quantity"])
        planned = float(row["planned_quantity"])
        if qualified > actual:
            raise ValueError("Qualified quantity cannot exceed actual quantity.")
        analysis = analyse_production(
            planned, actual, qualified, current_app.config["LOW_YIELD_THRESHOLD"]
        )
        cursor.execute(
            """
            UPDATE system_import.production_records
            SET qualified_quantity = %s, qualified_pending = FALSE,
                qualified_rate = %s, row_version = row_version + 1
            WHERE id = %s
            """,
            (qualified, qualified / actual if actual > 0 else None, record_id),
        )
        _audit(
            cursor,
            actor_id,
            "production_record",
            record_id,
            "qualified_quantity_follow_up",
            {"qualified_quantity": qualified},
        )
        if analysis.status == "open":
            cursor.execute(
                """
                INSERT INTO exceptions_queue (source_type, source_id, severity, summary, evidence_json)
                VALUES ('production', %s, %s, %s, %s)
                """,
                (record_id, analysis.severity, analysis.summary, json.dumps(analysis.evidence)),
            )
    return record_id


def create_abnormality(payload, actor_id=None):
    if uses_slitting_workflow():
        payload = slitting.normalise_abnormality(payload)
    payload = {**payload, "event_date": date.today().isoformat()}
    description = str(payload.get("description", "")).strip()
    immediate_action = str(payload.get("immediate_action", "")).strip()
    if not description:
        raise ValueError("Describe what happened.")
    payload["description"], payload["immediate_action"] = description, immediate_action
    duration = _downtime_minutes(payload.get("start_time"), payload.get("end_time"))
    cost_failure_types = payload.get("cost_failure_types", [])
    if not isinstance(cost_failure_types, list):
        raise ValueError("Cost-failure types must be a list.")
    is_resolved = payload.get("is_resolved", "no")
    effective_time_cost = payload.get("effective_time_cost") or None
    machine_type = payload.get("machine_type")
    if is_resolved not in {"yes", "no"}:
        raise ValueError("Is resolved must be yes or no.")
    responsible_person = str(payload.get("responsible_person", "")).strip()
    target_finish_date = payload.get("target_finish_date") or None
    if is_resolved == "yes":
        target_finish_date = date.today().isoformat()
    if is_resolved == "no" and (not responsible_person or not target_finish_date):
        raise ValueError("Unresolved events require a responsible person and expected finish date.")
    if is_resolved == "no" and not immediate_action:
        raise ValueError("Unresolved events require an immediate action.")
    if effective_time_cost not in {None, "yes", "no"}:
        raise ValueError("Effective time cost must be yes or no.")
    if machine_type not in {"semi", "auto"}:
        raise ValueError("Choose Semi / 半自动 or Auto / 全自动 machine type.")
    analysis = analyse_abnormality(
        payload["event_type"],
        payload["severity"],
        duration,
        payload.get("description", ""),
    )
    if duration < 0:
        raise ValueError("Duration cannot be negative.")
    workflow_state = "resolved" if is_resolved == "yes" else "open"

    client_record_id, payload_hash = _sync_identity(payload, "abnormality")
    with transaction() as cursor:
        _validate_active_event_type(cursor, payload)
        cursor.execute(
            "SELECT type_code FROM system_import.cost_failure_types WHERE is_active = TRUE"
        )
        active_types = {row["type_code"] for row in cursor.fetchall()}
        if not set(cost_failure_types).issubset(active_types):
            raise ValueError("An invalid or inactive cost-failure type was supplied.")
        replay_id = _existing_sync(cursor, "abnormality", client_record_id, payload_hash)
        if replay_id:
            return replay_id, analysis

        cursor.execute(
            """
            INSERT INTO system_import.abnormality_reports (
                event_date, start_time, end_time, shift_name, machine_code, machine_type, event_type, severity,
                duration_minutes, description, immediate_action, reported_by, custom_fields, state,
                is_resolved, effective_time_cost, cost_failure_types, responsible_person, target_finish_date
            ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
            """,
            (
                payload["event_date"],
                payload.get("start_time") or None,
                payload.get("end_time") or None,
                payload["shift_name"],
                payload["machine_code"],
                machine_type,
                payload["event_type"],
                payload["severity"],
                duration,
                payload.get("description", ""),
                payload.get("immediate_action", ""),
                actor_id,
                json.dumps(payload.get("custom_fields", {})),
                workflow_state,
                is_resolved,
                effective_time_cost,
                json.dumps(cost_failure_types),
                responsible_person or None,
                target_finish_date,
            ),
        )
        report_id = cursor.lastrowid
        cursor.execute(
            """
            INSERT INTO system_import.sync_receipts
                (source_type, client_record_id, payload_hash, server_record_id)
            VALUES ('abnormality', %s, %s, %s)
            """,
            (client_record_id, payload_hash, report_id),
        )
        _audit(cursor, actor_id, "abnormality_report", report_id, "created_kiosk", payload)
        if analysis.status == "open":
            cursor.execute(
                """
                INSERT INTO exceptions_queue
                    (source_type, source_id, severity, summary, evidence_json)
                VALUES ('abnormality', %s, %s, %s, %s)
                """,
                (report_id, analysis.severity, analysis.summary, json.dumps(analysis.evidence)),
            )
    return report_id, analysis


def list_records(
    record_type,
    manager=False,
    page=1,
    page_size=50,
    record_date=None,
    machine_code=None,
    formula_code=None,
    event_type=None,
):
    queries = {
        "production": """
            SELECT id, record_date, shift_name, machine_code, formula_code, batch_number,
                   planned_quantity, actual_quantity, qualified_quantity, qualified_pending, custom_fields,
                   achievement_rate, qualified_rate, state, created_at
            FROM system_import.production_records
        """,
        "abnormality": """
            SELECT id, event_date, shift_name, machine_code, event_type, severity,
                   responsible_person, duration_minutes, description, immediate_action,
                   custom_fields, state, created_at
            FROM system_import.abnormality_reports
        """,
        "tracker": """
            SELECT id, event_date, start_time, end_time, shift_name, machine_code, machine_type, event_type, severity,
                   duration_minutes, description, immediate_action, responsible_person, target_finish_date,
                   solution_provided, actual_finish_date, effectiveness, is_resolved, custom_fields, state
            FROM system_import.abnormality_reports
        """,
        "ongoing": """
            SELECT id, event_date, start_time, end_time, shift_name, machine_code, machine_type, event_type,
                   severity, is_resolved, effective_time_cost, solution_provided,
                   cost_failure_types, responsible_person, target_finish_date,
                   actual_finish_date, description, custom_fields, state, created_at
            FROM system_import.abnormality_reports
        """,
        "analysis": """
            SELECT source_type, source_id, severity, summary, evidence_json, status, created_at
            FROM exceptions_queue
        """,
    }
    table_map = {
        "production": "system_import.production_records",
        "abnormality": "system_import.abnormality_reports",
        "tracker": "system_import.abnormality_reports",
        "ongoing": "system_import.abnormality_reports",
        "analysis": "exceptions_queue",
    }
    order_by = {
        "production": "created_at DESC",
        "abnormality": "created_at DESC",
        "tracker": "event_date DESC, start_time DESC, created_at DESC",
        "ongoing": "event_date DESC, start_time DESC, created_at DESC",
        "analysis": "status ASC, created_at DESC",
    }
    conditions, params = [], []
    if record_type == "ongoing":
        conditions.append("is_resolved = 'no'")
    if record_type != "analysis" and record_date:
        conditions.append("record_date = %s" if record_type == "production" else "event_date = %s")
        params.append(record_date)
    if record_type != "analysis" and machine_code:
        conditions.append("machine_code = %s")
        params.append(machine_code)
    if record_type == "production" and formula_code:
        conditions.append("formula_code = %s")
        params.append(formula_code)
    if record_type in {"abnormality", "tracker", "ongoing"} and event_type:
        conditions.append("event_type = %s")
        params.append(event_type)
    where_clause = f" WHERE {' AND '.join(conditions)}" if conditions else ""
    offset = (page - 1) * page_size
    with transaction() as cursor:
        cursor.execute(
            f"SELECT COUNT(*) AS total FROM {table_map[record_type]}{where_clause}",
            params,
        )
        total = cursor.fetchone()["total"]
        cursor.execute(
            f"{queries[record_type]}{where_clause} ORDER BY {order_by[record_type]} LIMIT %s OFFSET %s",
            [*params, page_size, offset],
        )
        rows = cursor.fetchall()

    if record_type == "tracker" and not manager:
        for row in rows:
            row.pop("solution_provided", None)
    if uses_slitting_workflow():
        rows = [slitting.enrich(row) for row in rows]
    return _json_safe_rows(rows), total


DATA_EXPORT_COLUMNS = {
    "production": [
        ("id", "ID"), ("record_date", "Date / 日期"), ("shift_name", "Shift / 班次"),
        ("machine_code", "Machine / 设备"), ("formula_code", "Formula or model / 配方或型号"),
        ("batch_number", "Batch / 批次"), ("planned_quantity", "Planned quantity / 计划数量"),
        ("actual_quantity", "Actual quantity / 实际数量"),
        ("qualified_quantity", "Qualified quantity / 合格数量"),
        ("qualified_pending", "Qualification pending / 待补录合格"),
        ("achievement_rate", "Achievement rate / 达成率"),
        ("qualified_rate", "Qualified rate / 合格率"), ("notes", "Notes / 备注"),
        ("created_at", "Created / 创建时间"),
        ("planned_rolls", "Planned rolls / 计划卷数"), ("input_rolls", "Input rolls / 投入卷数"),
        ("qualified_rolls", "Qualified rolls / 合格卷数"), ("input_area", "Input area / 投入面积"),
        ("yield_rate", "Yield / 综收"),
    ],
    "abnormality": [
        ("id", "ID"), ("event_date", "Date / 日期"), ("start_time", "Start time / 开始时间"),
        ("end_time", "End time / 结束时间"), ("shift_name", "Shift / 班次"),
        ("machine_code", "Machine / 设备"), ("machine_type", "Process type / 工艺类型"),
        ("event_type", "Event type or main category / 事件类型或主类别"),
        ("subcategory", "Subcategory / 子类别"), ("severity", "Severity / 严重程度"),
        ("duration_minutes", "Duration minutes / 持续分钟"),
        ("description", "What happened / 发生了什么"),
        ("immediate_action", "Immediate action / 立即措施"),
        ("is_resolved", "Resolved / 已解决"), ("state", "Status / 状态"),
        ("responsible_person", "Responsible person / 责任人"),
        ("target_finish_date", "Expected finish / 预计完成"),
        ("actual_finish_date", "Actual finish / 实际完成"),
        ("solution_provided", "Solution / 解决方案"),
        ("effective_time_cost", "Effective time cost / 有效时间成本"),
        ("effectiveness", "Effectiveness / 有效性"), ("event_notes", "Event notes / 事件备注"),
        ("created_at", "Created / 创建时间"),
    ],
}


def data_export_columns(dataset):
    if dataset not in DATA_EXPORT_COLUMNS:
        raise ValueError("Choose Production or Abnormality data.")
    columns = DATA_EXPORT_COLUMNS[dataset]
    if dataset == "production" and not uses_slitting_workflow():
        columns = [item for item in columns if item[0] not in {
            "planned_rolls", "input_rolls", "qualified_rolls", "input_area", "yield_rate"
        }]
    if dataset == "abnormality" and not uses_slitting_workflow():
        columns = [item for item in columns if item[0] not in {"subcategory", "event_notes"}]
    return columns


def _export_date_range(start_date, end_date):
    try:
        start = date.fromisoformat(start_date)
        end = date.fromisoformat(end_date)
    except (TypeError, ValueError) as exc:
        raise ValueError("Choose valid export start and end dates.") from exc
    if start > end:
        raise ValueError("The export start date cannot be after the end date.")
    if (end - start).days > 366:
        raise ValueError("Choose an export range of one year or less.")
    return start, end


def data_export_rows(dataset, start_date, end_date):
    start, end = _export_date_range(start_date, end_date)
    if dataset == "production":
        query = """SELECT id, record_date, shift_name, machine_code, formula_code, batch_number,
                          planned_quantity, actual_quantity, qualified_quantity, qualified_pending,
                          achievement_rate, qualified_rate, notes, custom_fields, created_at
                   FROM system_import.production_records
                   WHERE record_date BETWEEN %s AND %s
                   ORDER BY record_date, shift_name, machine_code, id"""
    elif dataset == "abnormality":
        query = """SELECT id, event_date, start_time, end_time, shift_name, machine_code,
                          machine_type, event_type, severity, duration_minutes, description,
                          immediate_action, is_resolved, state, responsible_person,
                          target_finish_date, actual_finish_date, solution_provided,
                          effective_time_cost, effectiveness, custom_fields, created_at
                   FROM system_import.abnormality_reports
                   WHERE event_date BETWEEN %s AND %s
                   ORDER BY event_date, start_time, machine_code, id"""
    else:
        raise ValueError("Choose Production or Abnormality data.")
    with transaction() as cursor:
        cursor.execute(query, (start, end))
        rows = cursor.fetchall()
    if uses_slitting_workflow():
        rows = [slitting.enrich(row) for row in rows]
    return _json_safe_rows(rows)


def prepare_data_export(dataset, start_date, end_date, selected_columns=None):
    available = data_export_columns(dataset)
    available_keys = {key for key, _ in available}
    columns = list(selected_columns or [])
    if not columns:
        columns = [key for key, _ in available]
    if any(key not in available_keys for key in columns):
        raise ValueError("One or more selected export columns are invalid.")
    labels = dict(available)
    rows = data_export_rows(dataset, start_date, end_date)
    return columns, labels, rows


def export_data_workbook(dataset, start_date, end_date, selected_columns):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError as exc:
        raise RuntimeError("Excel export is unavailable. Install requirements.txt first.") from exc
    columns, labels, rows = prepare_data_export(dataset, start_date, end_date, selected_columns)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Production" if dataset == "production" else "Abnormalities"
    sheet.append([labels[key] for key in columns])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="164E70")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in rows:
        sheet.append([row.get(key, "") for key in columns])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, key in enumerate(columns, 1):
        values = [labels[key], *[str(row.get(key, "") or "") for row in rows[:200]]]
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = min(48, max(12,
            max(len(value) for value in values) + 2))
        for cell in sheet.iter_cols(min_col=index, max_col=index, min_row=2):
            for item in cell:
                item.alignment = Alignment(wrap_text=True, vertical="top")
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return output, len(rows)


def update_tracker(report_id, payload, manager=False):
    operator_fields = {
        "responsible_person",
        "target_finish_date",
        "actual_finish_date",
        "solution_provided",
    }
    manager_fields = operator_fields | {"solution_provided", "effectiveness", "state"}
    permitted_fields = manager_fields if manager else operator_fields
    changes = {key: value for key, value in payload.items() if key in permitted_fields}
    if not changes:
        raise ValueError("No permitted tracker fields supplied.")

    # Completion payloads from the close-event form must close the record even
    # when manager access is currently unlocked.
    has_finish = bool(str(changes.get("actual_finish_date", "")).strip())
    has_solution = bool(str(changes.get("solution_provided", "")).strip())
    if has_finish and has_solution and "state" not in changes:
        changes.update({"is_resolved": "yes", "state": "resolved"})

    if not manager:
        with transaction() as cursor:
            cursor.execute(
                "SELECT state, is_resolved FROM system_import.abnormality_reports WHERE id = %s",
                (report_id,),
            )
            row = cursor.fetchone()
            if not row or row["state"] != "open" or row["is_resolved"] != "no":
                raise ValueError(
                    "Only unresolved open events can be updated from the ongoing tracker."
                )
        if not (has_finish and has_solution):
            raise ValueError(
                "Enter both actual finish date and solution provided to close this event."
            )
        changes.update({"is_resolved": "yes", "state": "resolved"})

    if changes.get("state") == "resolved" and not changes.get("solution_provided"):
        raise ValueError("Manager resolution requires a documented solution.")

    assignments = ", ".join(f"{key} = %s" for key in changes)
    with transaction() as cursor:
        cursor.execute(
            f"""
            UPDATE system_import.abnormality_reports
            SET {assignments}, row_version = row_version + 1
            WHERE id = %s
            """,
            (*changes.values(), report_id),
        )
        if cursor.rowcount != 1:
            raise ValueError("Abnormality report not found.")
        _audit(cursor, None, "abnormality_report", report_id, "tracker_updated", changes)
    return {"closed": changes.get("state") == "resolved", "id": report_id}


def get_field_definitions(form_key, include_inactive=False):
    active_clause = "" if include_inactive else "AND is_active = TRUE"
    with transaction() as cursor:
        cursor.execute(
            f"""
            SELECT field_key, label, input_type, options_json, is_required, display_order
            FROM system_import.field_definitions
            WHERE form_key = %s {active_clause}
              AND NOT (form_key = 'production' AND field_key = 'work_order')
            ORDER BY display_order
            """,
            (form_key,),
        )
        return cursor.fetchall()


STANDARD_FIELD_CATALOG = {
    "production": [
        ("shift_name", "Shift / 班次", "select"),
        ("machine_code", "Machine / 设备", "select"),
        ("formula_code", "Formula / 配方", "select"),
        ("batch_number", "Batch number / 批次号", "text"),
        ("planned_quantity", "Planned quantity / 计划数量", "number"),
        ("actual_quantity", "Actual quantity / 实际数量", "number"),
        ("qualified_quantity", "Qualified quantity / 合格数量", "number"),
        ("notes", "Notes / 备注", "textarea"),
    ],
    "abnormality": [
        ("event_time", "Event time / 事件时间", "time_range"),
        ("shift_name", "Shift / 班次", "select"),
        ("machine_code", "Machine / 设备", "select"),
        ("machine_type", "Process type / 工艺类型", "select"),
        ("event_type", "Event type / 事件类型", "select"),
        ("severity", "Severity | Priority / 严重程度 | 优先级", "select"),
        ("is_resolved", "Is it resolved? / 是否已解决", "select"),
        ("responsible_person", "Responsible person / 责任人", "text"),
        ("target_finish_date", "Expected finish date / 预计完成日期", "date"),
        ("effective_time_cost", "是否为有效时间成本 / Effective time cost?", "select"),
        ("description", "What happened? / 发生了什么？", "textarea"),
        ("immediate_action", "Immediate action / 立即措施", "textarea"),
        ("actual_finish_date", "Actual finish date / 实际完成日期", "date"),
        ("solution_provided", "Solution provided / 解决方案", "textarea"),
    ],
}

SLITTING_STANDARD_FIELD_CATALOG = {
    "production": [
        ("shift_name", "Shift / 班次", "select"),
        ("machine_code", "Machine / 机台", "select"),
        ("slitting_model", "Model / 型号", "select"),
        ("planned_rolls", "Planned input rolls / 计划卷数", "number"),
        ("input_rolls", "Input rolls / 投入卷数", "number"),
        ("qualified_rolls", "Qualified output rolls / 合格产出卷数", "number"),
    ],
    "abnormality": [
        ("event_time", "Event time / 事件时间", "time_range"),
        ("shift_name", "Shift / 班次", "select"),
        ("machine_code", "Machine / 机台", "select"),
        ("event_type", "Main category / 主类别", "select"),
        ("subcategory", "Subcategory / 子类别", "select"),
        ("severity", "Severity | Priority / 严重程度 | 优先级", "select"),
        ("is_resolved", "Is it resolved? / 是否已解决", "select"),
        ("responsible_person", "Responsible person / 责任人", "text"),
        ("target_finish_date", "Expected finish date / 预计完成日期", "date"),
        ("effective_time_cost", "是否为有效时间成本 / Effective time cost?", "select"),
        ("event_notes", "Notes / 备注", "textarea"),
        ("immediate_action", "Immediate action / 立即措施", "textarea"),
        ("actual_finish_date", "Actual finish date / 实际完成日期", "date"),
        ("solution_provided", "Solution provided / 解决方案", "textarea"),
    ],
}

# Stored as value|bilingual display text so manager edits keep the database value stable.
STANDARD_FIELD_OPTIONS = {
    ("production", "shift_name"): ["Day|Day / 白班", "Night|Night / 夜班"],
    ("production", "machine_code"): [
        item
        for item in (
            "J1",
            "J2",
            "J3",
            "J4",
            "J5",
            "J6",
            "J7",
            "J8",
            "J9",
            "J10-1",
            "J10-2",
            "J10-3",
            "J10-4",
            "J10-5",
        )
    ],
    ("production", "formula_code"): ["K1-26", "B1-1", "E3", "custom|Custom / 自定义"],
    ("abnormality", "shift_name"): ["Day|Day / 白班", "Night|Night / 夜班"],
    ("abnormality", "machine_code"): [
        item
        for item in (
            "J1",
            "J2",
            "J3",
            "J4",
            "J5",
            "J6",
            "J7",
            "J8",
            "J9",
            "J10-1",
            "J10-2",
            "J10-3",
            "J10-4",
            "J10-5",
            "NA|NA / 不适用",
        )
    ],
    ("abnormality", "machine_type"): ["semi|Semi / 半自动", "auto|Auto / 全自动"],
    ("abnormality", "severity"): [
        "normal|Normal / 正常",
        "low|Low / 低",
        "medium|Medium / 中",
        "high|High / 高",
    ],
    ("abnormality", "is_resolved"): ["no|No / 否", "yes|Yes / 是"],
    ("abnormality", "effective_time_cost"): ["no|No / 否", "yes|Yes / 是"],
}


SLITTING_STANDARD_FIELD_OPTIONS = {
    ("production", "shift_name"): ["Day|Day / 白班", "Night|Night / 夜班"],
    ("production", "machine_code"): [item for item in slitting.MACHINES],
    ("abnormality", "shift_name"): ["Day|Day / 白班", "Night|Night / 夜班"],
    ("abnormality", "machine_code"): [*[item for item in slitting.MACHINES], "NA|NA / 不适用"],
    ("abnormality", "severity"): [
        "normal|Normal / 正常",
        "low|Low / 低",
        "medium|Medium / 中",
        "high|High / 高",
    ],
    ("abnormality", "is_resolved"): ["no|No / 否", "yes|Yes / 是"],
    ("abnormality", "effective_time_cost"): ["no|No / 否", "yes|Yes / 是"],
}


def _standard_field_catalog():
    return SLITTING_STANDARD_FIELD_CATALOG if uses_slitting_workflow() else STANDARD_FIELD_CATALOG


def _standard_field_options():
    return SLITTING_STANDARD_FIELD_OPTIONS if uses_slitting_workflow() else STANDARD_FIELD_OPTIONS


def get_standard_field_settings(form_key):
    catalog = _standard_field_catalog()
    option_defaults = _standard_field_options()
    if form_key not in catalog:
        raise ValueError("Unknown form.")
    with transaction() as cursor:
        cursor.execute(
            "SELECT field_key, label, help_text, options_json FROM system_import.standard_field_settings WHERE form_key = %s",
            (form_key,),
        )
        saved = {row["field_key"]: row for row in cursor.fetchall()}
    return [
        {
            "field_key": key,
            "label": saved.get(key, {}).get("label") or label,
            "help_text": saved.get(key, {}).get("help_text") or "",
            "input_type": input_type,
            "options_json": (
                saved.get(key, {}).get("options_json")
                if saved.get(key, {}).get("options_json") is not None
                else option_defaults.get((form_key, key), [])
            ),
            "is_configured": key in saved,
        }
        for key, label, input_type in catalog[form_key]
    ]


def save_standard_field_setting(form_key, payload):
    valid = {key for key, _, _ in _standard_field_catalog().get(form_key, [])}
    field_key = str(payload.get("field_key", "")).strip()
    label = str(payload.get("label", "")).strip()
    if field_key not in valid or not label:
        raise ValueError("Choose a standard field and provide a bilingual label.")
    help_text = str(payload.get("help_text", "")).strip()
    options = payload.get("options", [])
    if not isinstance(options, list):
        raise ValueError("Options must be a list.")
    with transaction() as cursor:
        cursor.execute(
            """
            INSERT INTO system_import.standard_field_settings
                (form_key, field_key, label, help_text, options_json)
            VALUES (%s, %s, %s, %s, %s)
            ON DUPLICATE KEY UPDATE label = VALUES(label), help_text = VALUES(help_text),
                options_json = VALUES(options_json)
            """,
            (
                form_key,
                field_key,
                label,
                help_text or None,
                json.dumps(options) if options else None,
            ),
        )
        _audit(
            cursor,
            None,
            "standard_field_setting",
            0,
            "saved",
            {"form_key": form_key, "field_key": field_key},
        )


def reset_standard_field_setting(form_key, field_key):
    valid = {key for key, _, _ in _standard_field_catalog().get(form_key, [])}
    if field_key not in valid:
        raise ValueError("Choose a valid standard field.")
    with transaction() as cursor:
        cursor.execute(
            "DELETE FROM system_import.standard_field_settings WHERE form_key=%s AND field_key=%s",
            (form_key, field_key),
        )


def get_cost_failure_types(include_inactive=False):
    query = "SELECT type_code, display_name, definition, is_active FROM system_import.cost_failure_types"
    if not include_inactive:
        query += " WHERE is_active = TRUE"
    query += " ORDER BY display_name"
    with transaction() as cursor:
        cursor.execute(query)
        return cursor.fetchall()


def get_event_types(include_inactive=False):
    if uses_slitting_workflow():
        return [
            {"event_value": key, "display_name": key, "is_active": True}
            for key in slitting.event_categories()
        ]
    query = "SELECT event_value, display_name, is_active FROM system_import.event_type_options"
    if not include_inactive:
        query += " WHERE is_active = TRUE"
    query += " ORDER BY display_order, display_name"
    with transaction() as cursor:
        cursor.execute(query)
        return cursor.fetchall()


def save_event_type(payload):
    if uses_slitting_workflow():
        raise ValueError("Slitting uses the approved main/subcategory catalog.")
    event_value = str(payload.get("event_value", "")).strip()
    display_name = str(payload.get("display_name", "")).strip()
    if not event_value or not display_name or len(event_value) > 80:
        raise ValueError("Event value and bilingual display name are required.")
    with transaction() as cursor:
        cursor.execute(
            """
            INSERT INTO system_import.event_type_options (event_value, display_name, display_order)
            VALUES (%s, %s, %s)
            ON DUPLICATE KEY UPDATE display_name = VALUES(display_name),
                display_order = VALUES(display_order), is_active = TRUE
            """,
            (event_value, display_name, int(payload.get("display_order", 100))),
        )
        _audit(cursor, None, "event_type_option", 0, "saved", {"event_value": event_value})


def deactivate_event_type(event_value):
    if uses_slitting_workflow():
        raise ValueError("Slitting uses the approved main/subcategory catalog.")
    with transaction() as cursor:
        cursor.execute(
            "UPDATE system_import.event_type_options SET is_active = FALSE WHERE event_value = %s",
            (event_value,),
        )
        if cursor.rowcount != 1:
            raise ValueError("Event type not found.")
        _audit(cursor, None, "event_type_option", 0, "deactivated", {"event_value": event_value})


def save_cost_failure_type(payload):
    type_code = str(payload.get("type_code", "")).strip().lower().replace(" ", "_")
    display_name = str(payload.get("display_name", "")).strip()
    definition = str(payload.get("definition", "")).strip()
    if not type_code.replace("_", "").isalnum() or not display_name or not definition:
        raise ValueError("Type code, display name, and definition are required.")
    with transaction() as cursor:
        cursor.execute(
            """
            INSERT INTO system_import.cost_failure_types (type_code, display_name, definition)
            VALUES (%s, %s, %s)
            ON DUPLICATE KEY UPDATE
                display_name = VALUES(display_name),
                definition = VALUES(definition),
                is_active = TRUE
            """,
            (type_code, display_name, definition),
        )
        _audit(cursor, None, "cost_failure_type", 0, "saved", {"type_code": type_code})


def deactivate_cost_failure_type(type_code):
    with transaction() as cursor:
        cursor.execute(
            "UPDATE system_import.cost_failure_types SET is_active = FALSE WHERE type_code = %s",
            (type_code,),
        )
        if cursor.rowcount != 1:
            raise ValueError("Cost-failure type not found.")
        _audit(cursor, None, "cost_failure_type", 0, "deactivated", {"type_code": type_code})


def save_field_definition(form_key, payload):
    valid_types = {"text", "number", "date", "select", "textarea"}
    if (
        form_key not in {"production", "abnormality"}
        or payload.get("input_type") not in valid_types
    ):
        raise ValueError("Invalid form or input type.")

    key = str(payload.get("field_key", "")).strip().lower().replace(" ", "_")
    if form_key == "production" and key == "work_order":
        raise ValueError("Work order is retired and cannot be added back to the production form.")
    if not key.replace("_", "").isalnum() or not key:
        raise ValueError("Field key may contain letters, numbers, and underscores only.")

    options = payload.get("options", [])
    with transaction() as cursor:
        cursor.execute(
            """
            INSERT INTO system_import.field_definitions (
                form_key, field_key, label, input_type, options_json, is_required, display_order
            ) VALUES (%s, %s, %s, %s, %s, %s, %s)
            ON DUPLICATE KEY UPDATE
                label = VALUES(label),
                input_type = VALUES(input_type),
                options_json = VALUES(options_json),
                is_required = VALUES(is_required),
                display_order = VALUES(display_order),
                is_active = TRUE
            """,
            (
                form_key,
                key,
                payload.get("label", key),
                payload["input_type"],
                json.dumps(options),
                bool(payload.get("is_required", False)),
                int(payload.get("display_order", 200)),
            ),
        )
        _audit(
            cursor, None, "field_definition", 0, "saved", {"form_key": form_key, "field_key": key}
        )


def manager_dashboard():
    with transaction() as cursor:
        cursor.execute("""
            SELECT COUNT(*) AS records, COALESCE(SUM(actual_quantity), 0) AS actual,
                   COALESCE(SUM(qualified_quantity), 0) AS qualified
            FROM system_import.production_records
            WHERE record_date = CURDATE()
            """)
        production = cursor.fetchone()
        cursor.execute("""
            SELECT COUNT(*) AS open_events, COALESCE(SUM(duration_minutes), 0) AS downtime
            FROM system_import.abnormality_reports
            WHERE event_date = CURDATE() AND state = 'open'
            """)
        abnormality = cursor.fetchone()
        cursor.execute("""
            SELECT machine_code, COUNT(*) AS event_count,
                   COALESCE(SUM(duration_minutes), 0) AS downtime
            FROM system_import.abnormality_reports
            WHERE event_date >= CURDATE() - INTERVAL 30 DAY
            GROUP BY machine_code
            ORDER BY downtime DESC LIMIT 5
            """)
        machines = cursor.fetchall()

    actual = float(production["actual"] or 0)
    qualified = float(production["qualified"] or 0)
    return {
        "today_records": production["records"],
        "today_yield": round(qualified / actual * 100, 1) if actual else None,
        "open_events": abnormality["open_events"],
        "downtime_minutes": int(abnormality["downtime"] or 0),
        "machines": machines,
    }


DAILY_TASK_TYPES = {
    "production",
    "cleaning",
    "custom",
    "slitting_production",
    "slitting_schedule",
    "slitting_training",
}
DAILY_TASK_MACHINES = {
    *(f"J{number}" for number in range(1, 10)),
    *(f"J10-{number}" for number in range(1, 6)),
}
SLITTING_TASK_PRODUCTS = {
    "513MM",
    "514.9MM",
    "515.6MM",
    "547MM",
    "242MM",
    "515.4MM",
    "270.6MM",
    "270MM",
    "88.8MM",
    "67.1MM",
    "73.8MM",
    "96MM",
    "264.5MM",
    "269.8MM",
    "61MM",
    "73.7MM",
    "196MM",
    "557.20MM",
    "62.3MM",
    "61.1MM",
    "61.6MM",
    "93.7MM",
    "73.5MM",
}
SLITTING_TASK_STATUSES = {"production_in_progress", "re_cut", "label_changing"}
SLITTING_TARGET_UNITS = {"sr", "MR", "custom"}

DISPLAY_BLOCK_DEFAULTS = [
    {
        "key": "task",
        "title": "Today's task / 今日任务",
        "visible": True,
        "order": 10,
        "width": 2,
        "height": 2,
        "font_size": 16,
    },
    {
        "key": "reminders",
        "title": "Reminders / 提醒事项",
        "visible": True,
        "order": 15,
        "width": 2,
        "height": 1,
        "font_size": 16,
    },
    {
        "key": "rates",
        "title": "Production rates / 生产率",
        "visible": True,
        "order": 20,
        "width": 1,
        "height": 1,
        "font_size": 16,
    },
    {
        "key": "pending_qualification",
        "title": "Pending qualification data / 待补录合格数据",
        "visible": True,
        "order": 25,
        "width": 1,
        "height": 1,
        "font_size": 16,
    },
    {
        "key": "production_trend",
        "title": "7-day production / 7日生产趋势",
        "visible": True,
        "order": 30,
        "width": 3,
        "height": 2,
        "font_size": 15,
    },
    {
        "key": "abnormality_trend",
        "title": "Abnormalities by date / 每日异常",
        "visible": True,
        "order": 35,
        "width": 2,
        "height": 2,
        "font_size": 15,
    },
    {
        "key": "ongoing",
        "title": "Ongoing abnormalities / 进行中异常",
        "visible": True,
        "order": 40,
        "width": 5,
        "height": 2,
        "font_size": 16,
    },
]
SLITTING_DISPLAY_BLOCK_DEFAULTS = [
    {
        "key": "slitting_task",
        "title": "Slitting task overview / 分切及包装任务概览",
        "visible": True,
        "order": 5,
        "width": 3,
        "height": 3,
        "font_size": 14,
    },
    {
        "key": "slitting_training",
        "title": "Training / 培训",
        "visible": True,
        "order": 10,
        "width": 2,
        "height": 2,
        "font_size": 14,
    },
    {
        "key": "slitting_reminders",
        "title": "Reminders / 提醒事项",
        "visible": True,
        "order": 15,
        "width": 2,
        "height": 1,
        "font_size": 15,
    },
    {
        "key": "slitting_pending_qualification",
        "title": "Pending qualified rolls / 待补录合格卷数",
        "visible": True,
        "order": 18,
        "width": 2,
        "height": 2,
        "font_size": 15,
    },
    {
        "key": "slitting_events",
        "title": "Ongoing abnormalities / 进行中异常",
        "visible": True,
        "order": 20,
        "width": 5,
        "height": 2,
        "font_size": 16,
    },
    {
        "key": "slitting_production_trend",
        "title": "7-day production by machine and width / 7日机台与幅宽生产",
        "visible": True,
        "order": 30,
        "width": 3,
        "height": 2,
        "font_size": 14,
    },
    {
        "key": "slitting_abnormality_trend",
        "title": "Abnormal downtime by date / 每日异常停机",
        "visible": True,
        "order": 35,
        "width": 2,
        "height": 2,
        "font_size": 14,
    },
]
DISPLAY_EVENT_COLUMN_DEFAULTS = [
    {"key": "event_date", "label": "Event date / 事件日期", "visible": True, "order": 10},
    {"key": "start_time", "label": "Start time / 开始时间", "visible": True, "order": 20},
    {"key": "end_time", "label": "End time / 结束时间", "visible": True, "order": 30},
    {"key": "shift_name", "label": "Shift / 班次", "visible": True, "order": 40},
    {"key": "machine_code", "label": "Machine / 设备", "visible": True, "order": 50},
    {"key": "machine_type", "label": "Process / 工艺", "visible": True, "order": 60},
    {"key": "event_type", "label": "Event type / 事件类型", "visible": True, "order": 70},
    {
        "key": "severity",
        "label": "Severity / Priority / 严重程度 / 优先级",
        "visible": True,
        "order": 80,
    },
    {
        "key": "duration_minutes",
        "label": "Duration (minutes) / 持续时间（分钟）",
        "visible": True,
        "order": 90,
    },
    {"key": "is_resolved", "label": "Resolved? / 是否解决", "visible": True, "order": 100},
    {
        "key": "responsible_person",
        "label": "Responsible person / 责任人",
        "visible": True,
        "order": 110,
    },
    {
        "key": "target_finish_date",
        "label": "Expected finish / 预计完成",
        "visible": True,
        "order": 120,
    },
    {
        "key": "effective_time_cost",
        "label": "Effective time cost? / 是否为有效时间成本",
        "visible": True,
        "order": 130,
    },
    {
        "key": "cost_failure_types",
        "label": "Cost-failure types / 成本失效类型",
        "visible": True,
        "order": 140,
    },
    {"key": "potential_cost", "label": "Potential cost / 潜在成本", "visible": True, "order": 150},
    {"key": "description", "label": "What happened? / 发生了什么？", "visible": True, "order": 160},
    {
        "key": "immediate_action",
        "label": "Immediate action / 立即措施",
        "visible": True,
        "order": 170,
    },
]

DISPLAY_BRANDING_DEFAULTS = {
    "title": "GNEM Slurry Production Details / GNEM 制浆生产明细",
    "subtitle": "Central factory records · Local SQL",
    "display_eyebrow": "GNEM SLURRY PRODUCTION DETAILS / GNEM 制浆生产明细",
    "display_title": "Daily Factory Display / 每日工厂看板",
}
SLITTING_DISPLAY_BRANDING_DEFAULTS = {
    "title": "GNEM Slitting Production / GNEM 分切生产",
    "subtitle": "Central factory records · Local SQL",
    "display_eyebrow": "GNEM SLITTING PRODUCTION / GNEM 分切生产",
    "display_title": "Daily Slitting Display / 分切每日看板",
}

SLITTING_EVENT_COLUMN_DEFAULTS = [
    {"key": "shift_name", "label": "Shift / \u73ed\u6b21", "visible": True, "order": 5},
    {"key": "machine_code", "label": "Machine / 机台", "visible": True, "order": 10},
    {"key": "event_type", "label": "Main category / 主类别", "visible": True, "order": 20},
    {"key": "subcategory", "label": "Subcategory / 子类别", "visible": True, "order": 30},
    {"key": "event_notes", "label": "Notes / 备注", "visible": True, "order": 40},
    {
        "key": "duration_minutes",
        "label": "Time loss (min) / 损失分钟",
        "visible": True,
        "order": 50,
    },
    {"key": "is_resolved", "label": "Resolved / 已解决", "visible": True, "order": 60},
]


def default_display_settings():
    if uses_slitting_workflow():
        branding = dict(SLITTING_DISPLAY_BRANDING_DEFAULTS)
        blocks = [dict(x) for x in SLITTING_DISPLAY_BLOCK_DEFAULTS]
        if current_department() == "coating":
            branding = {
                key: value.replace("SLITTING", "COATING")
                .replace("Slitting", "Coating")
                .replace("分切", "涂布")
                for key, value in branding.items()
            }
            for block in blocks:
                block["title"] = (
                    block["title"].replace("Slitting", "Coating").replace("分切", "涂布")
                )
        return {
            "layout_version": 4,
            "branding": branding,
            "blocks": blocks,
            "event_columns": [dict(x) for x in SLITTING_EVENT_COLUMN_DEFAULTS],
            "abnormality_chart_types": [],
            "abnormality_chart_types_configured": False,
        }
    return {
        "layout_version": 4,
        "branding": dict(DISPLAY_BRANDING_DEFAULTS),
        "blocks": [dict(x) for x in DISPLAY_BLOCK_DEFAULTS],
        "event_columns": [],
        "abnormality_chart_types": [],
        "abnormality_chart_types_configured": False,
    }


def _default_abnormality_chart_types():
    """Keep the historical chart behavior until a manager saves a selection."""
    excluded = {"cleaning", "normal production"}
    return [
        row["event_value"]
        for row in get_event_types()
        if str(row.get("event_value", "")).strip().lower() not in excluded
    ]


def get_display_settings():
    with transaction() as cursor:
        cursor.execute(
            "SELECT settings_json FROM system_import.display_settings WHERE setting_key = 'factory_display'"
        )
        row = cursor.fetchone()
    if not row:
        defaults = default_display_settings()
        defaults["abnormality_chart_types"] = _default_abnormality_chart_types()
        return defaults
    saved = (
        json.loads(row["settings_json"])
        if isinstance(row["settings_json"], str)
        else row["settings_json"]
    )
    if int(saved.get("layout_version", 1)) < 2:
        for block in saved.get("blocks", []):
            block["width"] = {1: 2, 2: 4, 3: 5}.get(int(block.get("width", 1)), 2)
            if block.get("type") == "static_text":
                block.setdefault("heading_level", "h2")
        saved["layout_version"] = 2
    if uses_slitting_workflow() and int(saved.get("layout_version", 1)) < 3:
        for block in saved.get("blocks", []):
            if block.get("key") == "slitting_task":
                block.update({"width": 3, "height": 3, "font_size": 14})
        saved["layout_version"] = 3
    if uses_slitting_workflow() and int(saved.get("layout_version", 1)) < 4:
        saved["blocks"] = [
            block for block in saved.get("blocks", []) if block.get("key") != "slitting_training"
        ]
        saved["layout_version"] = 4
    if current_department() == "slurry" and int(saved.get("layout_version", 1)) < 3:
        saved["blocks"] = [
            block
            for block in saved.get("blocks", [])
            if block.get("key") not in {"events", "reminders"}
        ]
        saved["event_columns"] = []
        saved["layout_version"] = 3
    defaults = default_display_settings()
    if isinstance(saved.get("branding"), dict):
        defaults["branding"].update(
            {
                key: str(saved["branding"].get(key, value))
                for key, value in defaults["branding"].items()
            }
        )
    for group in ("blocks", "event_columns"):
        saved_items = saved.get(group, [])
        saved_by_key = {item["key"]: item for item in saved_items}
        default_keys = {item["key"] for item in defaults[group]}
        defaults[group] = [
            {**item, **saved_by_key.get(item["key"], {})} for item in defaults[group]
        ]
        if group == "blocks":
            defaults[group].extend(
                item
                for item in saved_items
                if item.get("key") not in default_keys and item.get("type") == "static_text"
            )
    configured = bool(saved.get("abnormality_chart_types_configured", False))
    defaults["abnormality_chart_types_configured"] = configured
    if configured:
        defaults["abnormality_chart_types"] = [
            str(value).strip()
            for value in saved.get("abnormality_chart_types", [])
            if str(value).strip()
        ]
    else:
        defaults["abnormality_chart_types"] = _default_abnormality_chart_types()
    if current_department() == "coating":
        replacements = (
            ("SLITTING", "COATING"),
            ("Slitting", "Coating"),
            ("slitting", "coating"),
            ("分切", "涂布"),
        )

        def coating_label(value):
            value = str(value)
            for source, target in replacements:
                value = value.replace(source, target)
            return value

        defaults["branding"] = {
            key: coating_label(value) for key, value in defaults["branding"].items()
        }
        for block in defaults["blocks"]:
            block["title"] = coating_label(block["title"])
    return defaults


def save_display_settings(payload, actor_id=None):
    defaults = default_display_settings()
    cleaned = {"layout_version": defaults["layout_version"]}
    branding = payload.get("branding", defaults["branding"])
    if not isinstance(branding, dict):
        raise ValueError("Branding must be an object.")
    cleaned_branding = {
        key: str(branding.get(key, fallback)).strip()
        for key, fallback in defaults["branding"].items()
    }
    if (
        not cleaned_branding["title"]
        or not cleaned_branding["display_title"]
        or any(len(value) > 160 for value in cleaned_branding.values())
    ):
        raise ValueError(
            "Application and display titles are required and header text must be 160 characters or fewer."
        )
    cleaned["branding"] = cleaned_branding
    selected_types = payload.get("abnormality_chart_types")
    if not isinstance(selected_types, list):
        raise ValueError("Abnormality chart event types must be a list.")
    cleaned["abnormality_chart_types"] = list(
        dict.fromkeys(
            value
            for value in (str(item).strip() for item in selected_types)
            if value and len(value) <= 160
        )
    )
    cleaned["abnormality_chart_types_configured"] = True
    for group in ("blocks", "event_columns"):
        if not isinstance(payload.get(group), list):
            raise ValueError(f"{group} must be a list.")
        supplied = {str(x.get("key")): x for x in payload[group] if isinstance(x, dict)}
        items = []
        for fallback in defaults[group]:
            item, label_key = supplied.get(fallback["key"], {}), (
                "title" if group == "blocks" else "label"
            )
            label = str(item.get(label_key, fallback[label_key])).strip()
            if not label or len(label) > 160:
                raise ValueError("Display titles must contain 1 to 160 characters.")
            try:
                order = max(0, min(int(item.get("order", fallback["order"])), 999))
            except (TypeError, ValueError) as exc:
                raise ValueError("Display order must be a whole number.") from exc
            result = {
                "key": fallback["key"],
                label_key: label,
                "visible": bool(item.get("visible", False)),
                "order": order,
            }
            if group == "blocks":
                result["width"] = max(1, min(int(item.get("width", fallback["width"])), 5))
                result["height"] = max(
                    1, min(int(item.get("height", fallback.get("height", 1))), 5)
                )
                result["font_size"] = max(
                    12, min(int(item.get("font_size", fallback.get("font_size", 16))), 36)
                )
            items.append(result)
        if group == "blocks":
            for item in payload["blocks"]:
                key = str(item.get("key", ""))
                if (
                    key in {fallback["key"] for fallback in defaults["blocks"]}
                    or item.get("type") != "static_text"
                ):
                    continue
                if not re.fullmatch(r"text_[a-f0-9-]{8,40}", key):
                    raise ValueError("Permanent text block identifier is invalid.")
                title, content = (
                    str(item.get("title", "")).strip(),
                    str(item.get("content", "")).strip(),
                )
                if not title or len(title) > 160:
                    raise ValueError("Permanent block titles must contain 1 to 160 characters.")
                if len(content) > 4000:
                    raise ValueError("Permanent text must be 4,000 characters or fewer.")
                heading_level = str(item.get("heading_level", "h2")).lower()
                if heading_level not in {"h1", "h2"}:
                    raise ValueError("Permanent block heading must be Heading 1 or Heading 2.")
                items.append(
                    {
                        "key": key,
                        "type": "static_text",
                        "title": title,
                        "content": content,
                        "heading_level": heading_level,
                        "visible": bool(item.get("visible", True)),
                        "order": max(0, min(int(item.get("order", 900)), 999)),
                        "width": max(1, min(int(item.get("width", 5)), 5)),
                        "height": max(1, min(int(item.get("height", 1)), 5)),
                        "font_size": max(12, min(int(item.get("font_size", 16)), 36)),
                    }
                )
        cleaned[group] = items
    with transaction() as cursor:
        cursor.execute(
            "INSERT INTO system_import.display_settings (setting_key, settings_json) VALUES ('factory_display', %s) ON DUPLICATE KEY UPDATE settings_json = VALUES(settings_json)",
            (json.dumps(cleaned),),
        )
        _audit(cursor, actor_id, "display_settings", 0, "saved", cleaned)
    return cleaned


def _legacy_task_items(task):
    """Map the former one-row task format to repeatable items without losing data."""
    types = json.loads(task.get("task_types") or "[]")
    items = []
    for task_type in types:
        item = {"type": task_type, "description": ""}
        if task_type == "production":
            item.update(
                {
                    "formula_code": task.get("formula_code") or "",
                    "amount_needed": task.get("amount_needed"),
                    "machine_assigned": task.get("machine_assigned") or "",
                }
            )
        elif task_type == "custom":
            item["description"] = task.get("custom_task") or ""
        items.append(item)
    return items


def current_production_slot(now=None):
    """Return the active production date and department-specific shift."""
    now = now or datetime.now()
    if current_department() == "slurry":
        production_date = now.date()
        if production_date.weekday() == 6:
            production_date += timedelta(days=1)
        return production_date, "Day 08-20"
    shift_minutes = now.hour * 60 + now.minute
    if shift_minutes < 7 * 60 + 50:
        return now.date() - timedelta(days=1), "Night"
    if shift_minutes >= 19 * 60 + 50:
        return now.date(), "Night"
    return now.date(), "Day"


def production_shift_label(production_date, shift_name):
    """Return the scheduled hours shown beside a task slot."""
    if isinstance(production_date, str):
        production_date = date.fromisoformat(production_date)
    if current_department() == "slurry":
        if shift_name == "Day 08-17":
            hours = "08:00–17:00"
        else:
            hours = "08:00–20:00"
        return f"Day · {hours}"
    return "Day · 07:50–19:50" if shift_name == "Day" else "Night · 19:50–07:50"


def get_daily_task(record_date=None, shift_name=None):
    """Return one shift-specific task; absent slots are explicit."""
    current_date, current_shift = current_production_slot()
    selected_date = record_date or current_date
    selected_shift = str(shift_name or current_shift).title()
    allowed_shifts = (
        {"Day", "Night"} if uses_slitting_workflow() else {"Day", "Day 08-20", "Day 08-17"}
    )
    if selected_shift not in allowed_shifts:
        raise ValueError("Choose Day or Night shift.")
    if isinstance(selected_date, str):
        try:
            selected_date = date.fromisoformat(selected_date)
        except ValueError as exc:
            raise ValueError("Choose a valid date.") from exc
    with transaction() as cursor:
        cursor.execute(
            """
            SELECT record_date, shift_name, task_types, task_items, reminders, custom_task, formula_code, amount_needed, machine_assigned, updated_at
            FROM system_import.daily_tasks
            WHERE record_date = %s AND shift_name = %s
            """,
            (selected_date, selected_shift),
        )
        task = cursor.fetchone()
    if not task:
        return {
            "record_date": selected_date.isoformat(),
            "shift_name": selected_shift,
            "shift_label": production_shift_label(selected_date, selected_shift),
            "task_items": [],
            "reminders": [],
            "updated_at": None,
        }
    task["task_items"] = json.loads(task["task_items"] or "null") or _legacy_task_items(task)
    task["reminders"] = json.loads(task["reminders"] or "[]")
    task["shift_label"] = production_shift_label(selected_date, selected_shift)
    return _json_safe_rows([task])[0]


def save_daily_task(payload, actor_id=None):
    """Save an employee-entered task for a production date and shift."""
    default_date, default_shift = current_production_slot()
    try:
        task_date = date.fromisoformat(str(payload.get("record_date") or default_date))
    except ValueError as exc:
        raise ValueError("Choose a valid task date.") from exc
    shift_name = str(payload.get("shift_name") or default_shift).title()
    allowed_shifts = (
        {"Day", "Night"} if uses_slitting_workflow() else {"Day", "Day 08-20", "Day 08-17"}
    )
    if shift_name not in allowed_shifts:
        raise ValueError("Choose Day or Night shift.")
    raw_items = payload.get("task_items", [])
    if not isinstance(raw_items, list) or not raw_items or len(raw_items) > 20:
        raise ValueError("Add between one and twenty task items.")
    raw_reminders = payload.get("reminders", [])
    if not isinstance(raw_reminders, list) or len(raw_reminders) > 12:
        raise ValueError("Add up to twelve reminder items.")
    reminders = []
    for reminder in raw_reminders:
        clean_reminder = str(reminder).strip()
        if clean_reminder:
            if len(clean_reminder) > 300:
                raise ValueError("Each reminder must be 300 characters or fewer.")
            reminders.append(clean_reminder)

    task_items = []
    if uses_slitting_workflow():
        production_count = schedule_count = training_count = 0
        for raw_item in raw_items:
            if not isinstance(raw_item, dict):
                raise ValueError("Each task item must be valid.")
            task_type = str(raw_item.get("type", "")).strip().lower()
            if task_type == "slitting_production":
                production_count += 1
                product_item = (
                    str(raw_item.get("product_item", "")).replace(" ", "").strip().upper()
                )
                custom_product_item = str(raw_item.get("custom_product_item", "")).strip()
                status = str(raw_item.get("production_status", "")).strip()
                custom_status = str(raw_item.get("custom_status", "")).strip()
                target_unit = str(raw_item.get("target_unit", "")).strip()
                custom_unit = str(raw_item.get("custom_unit", "")).strip()
                if product_item == "CUSTOM":
                    if not custom_product_item or len(custom_product_item) > 120:
                        raise ValueError("Enter a custom product item of 120 characters or fewer.")
                    product_item = custom_product_item
                elif product_item not in SLITTING_TASK_PRODUCTS:
                    raise ValueError("Choose a listed Slitting production item or Custom.")
                if status == "CUSTOM":
                    if not custom_status or len(custom_status) > 120:
                        raise ValueError(
                            "Enter a custom production status of 120 characters or fewer."
                        )
                    status = custom_status
                elif status not in SLITTING_TASK_STATUSES:
                    raise ValueError("Choose a valid Slitting production status or Custom.")
                if target_unit not in SLITTING_TARGET_UNITS:
                    raise ValueError(
                        "Specify whether the production target uses sr, MR, or a custom unit."
                    )
                if target_unit == "custom" and (not custom_unit or len(custom_unit) > 30):
                    raise ValueError(
                        "Enter a custom production target unit of 30 characters or fewer."
                    )
                try:
                    target = float(raw_item.get("production_target"))
                except (TypeError, ValueError) as exc:
                    raise ValueError("Production target must be a number.") from exc
                if target < 0 or target > 1000000:
                    raise ValueError("Production target cannot be negative.")
                task_items.append(
                    {
                        "type": task_type,
                        "product_item": product_item,
                        "production_status": status,
                        "production_target": target,
                        "target_unit": target_unit,
                        "custom_unit": custom_unit if target_unit == "custom" else "",
                    }
                )
            elif task_type == "slitting_schedule":
                schedule_count += 1
                machine = str(raw_item.get("machine", "")).strip().upper()
                product_item = (
                    str(raw_item.get("product_item", "")).replace(" ", "").strip().upper()
                )
                custom_product_item = str(raw_item.get("custom_product_item", "")).strip()
                planned_input = str(raw_item.get("planned_input", "")).strip()
                operator = str(raw_item.get("operator", "")).strip()
                shift_lead = str(raw_item.get("shift_lead", "")).strip()
                if machine not in {
                    choice["value"] for choice in slitting.machine_choices("production")
                }:
                    raise ValueError("Choose a configured Slitting machine.")
                if product_item == "CUSTOM":
                    if not custom_product_item or len(custom_product_item) > 120:
                        raise ValueError(
                            "Enter a custom scheduled product item of 120 characters or fewer."
                        )
                    product_item = custom_product_item
                elif product_item not in SLITTING_TASK_PRODUCTS:
                    raise ValueError("Choose a listed Slitting production item or Custom.")
                if not planned_input or len(planned_input) > 120:
                    raise ValueError(
                        "Schedule planned input is required and must be 120 characters or fewer."
                    )
                if not operator or len(operator) > 120:
                    raise ValueError(
                        "Schedule operator is required and must be 120 characters or fewer."
                    )
                if not shift_lead or len(shift_lead) > 120:
                    raise ValueError(
                        "Schedule shift lead is required and must be 120 characters or fewer."
                    )
                task_items.append(
                    {
                        "type": task_type,
                        "machine": machine,
                        "product_item": product_item,
                        "planned_input": planned_input,
                        "operator": operator,
                        "shift_lead": shift_lead,
                    }
                )
            elif task_type == "slitting_training":
                training_count += 1
                training = {
                    key: str(raw_item.get(key, "")).strip()
                    for key in ("training_item", "training_type", "subject", "time")
                }
                if any(not value for value in training.values()):
                    raise ValueError(
                        "Complete training item, type, person or group being trained, and time."
                    )
                if any(len(value) > 120 for value in training.values()):
                    raise ValueError("Each training field must be 120 characters or fewer.")
                task_items.append({"type": task_type, **training})
            else:
                raise ValueError("Slitting tasks must be production, schedule, or training items.")
        if not production_count and not training_count:
            raise ValueError("Add at least one Slitting production or training item.")
        if production_count and not schedule_count:
            raise ValueError("Add at least one machine schedule row for production.")
        if schedule_count and not production_count:
            raise ValueError("A machine schedule requires at least one production item.")
    else:
        task_items = _clean_standard_daily_task_items(raw_items)

    task_types = [item["type"] for item in task_items]
    first_production = next((item for item in task_items if item["type"] == "production"), {})
    first_custom = next((item for item in task_items if item["type"] == "custom"), {})

    with transaction() as cursor:
        cursor.execute(
            """
            INSERT INTO system_import.daily_tasks
                (record_date, shift_name, task_types, task_items, reminders, custom_task, formula_code, amount_needed, machine_assigned, updated_by)
            VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
            ON DUPLICATE KEY UPDATE
                task_types = VALUES(task_types), task_items = VALUES(task_items), reminders = VALUES(reminders), custom_task = VALUES(custom_task),
                formula_code = VALUES(formula_code), amount_needed = VALUES(amount_needed),
                machine_assigned = VALUES(machine_assigned), updated_by = VALUES(updated_by)
            """,
            (
                task_date,
                shift_name,
                json.dumps(task_types),
                json.dumps(task_items),
                json.dumps(reminders),
                first_custom.get("description") or None,
                first_production.get("formula_code"),
                first_production.get("amount_needed"),
                first_production.get("machine_assigned"),
                actor_id,
            ),
        )
        _audit(
            cursor,
            actor_id,
            "daily_task",
            0,
            "saved",
            {
                "record_date": task_date.isoformat(),
                "shift_name": shift_name,
                "task_types": task_types,
                "item_count": len(task_items),
                "reminder_count": len(reminders),
            },
        )


def _clean_standard_daily_task_items(raw_items):
    task_items = []
    for raw_item in raw_items:
        if not isinstance(raw_item, dict):
            raise ValueError("Each task item must be valid.")
        task_type = str(raw_item.get("type", "")).strip().lower()
        if task_type not in {"production", "cleaning", "custom"}:
            raise ValueError("Choose a valid task type for every item.")
        item = {"type": task_type, "description": str(raw_item.get("description", "")).strip()}
        if task_type == "custom" and not item["description"]:
            raise ValueError("Custom tasks require a description.")
        if task_type == "cleaning":
            process_type = str(raw_item.get("process_type", "")).strip().lower()
            if process_type not in {"semi", "auto"}:
                raise ValueError("Each cleaning task requires Semi / 半自动 or Auto / 全自动.")
            item["process_type"] = process_type
        if task_type == "production":
            formula_code = str(raw_item.get("formula_code", "")).strip().upper()
            machine_assigned = str(raw_item.get("machine_assigned", "")).strip().upper()
            if not formula_code:
                raise ValueError("Each production task requires a formula.")
            if machine_assigned not in DAILY_TASK_MACHINES:
                raise ValueError("Choose a machine for every production task.")
            try:
                amount_needed = float(raw_item.get("amount_needed"))
            except (TypeError, ValueError) as exc:
                raise ValueError("Each production task requires a valid tank amount.") from exc
            if amount_needed < 0:
                raise ValueError("Tank amount cannot be negative.")
            try:
                weight_needed = float(raw_item.get("weight_needed", 0) or 0)
            except (TypeError, ValueError) as exc:
                raise ValueError("Weight needed must be a valid number.") from exc
            if weight_needed < 0:
                raise ValueError("Weight needed cannot be negative.")
            item.update(
                {
                    "formula_code": formula_code,
                    "amount_needed": amount_needed,
                    "weight_needed": weight_needed,
                    "machine_assigned": machine_assigned,
                }
            )
        task_items.append(item)

    return task_items


def daily_report(record_date):
    try:
        selected_date = date.fromisoformat(record_date)
    except (TypeError, ValueError) as exc:
        raise ValueError("Choose a valid report date.") from exc

    with transaction() as cursor:
        cursor.execute(
            """
            SELECT record_date, shift_name, machine_code, formula_code, batch_number,
                   planned_quantity, actual_quantity, qualified_quantity,
                   achievement_rate, qualified_rate, notes, custom_fields, created_at
            FROM system_import.production_records
            WHERE record_date = %s
            ORDER BY shift_name, machine_code, batch_number
            """,
            (selected_date,),
        )
        production = cursor.fetchall()
        cursor.execute(
            """
            SELECT event_date, shift_name, machine_code, machine_type,
                   CONCAT(COALESCE(TIME_FORMAT(start_time, '%H:%i'), '—'), ' - ',
                          COALESCE(TIME_FORMAT(end_time, '%H:%i'), '—')) AS event_time,
                   duration_minutes, event_type, severity, description, immediate_action,
                   is_resolved, responsible_person, target_finish_date, actual_finish_date,
                   solution_provided, effectiveness, custom_fields
            FROM system_import.abnormality_reports
            WHERE event_date = %s
            ORDER BY start_time, machine_code
            """,
            (selected_date,),
        )
        events = cursor.fetchall()
        cursor.execute(
            """
            SELECT id, event_date, shift_name, machine_code, machine_type,
                   CONCAT(COALESCE(TIME_FORMAT(start_time, '%H:%i'), '—'), ' - ',
                          COALESCE(TIME_FORMAT(end_time, '%H:%i'), '—')) AS event_time,
                   duration_minutes, event_type, severity, description, immediate_action,
                   responsible_person, target_finish_date, custom_fields
            FROM system_import.abnormality_reports
            WHERE is_resolved = 'no'
            ORDER BY event_date, start_time, machine_code
            """
        )
        ongoing_events = cursor.fetchall()
        abnormality_trends = []
        production_trends = []
        if uses_slitting_workflow():
            cursor.execute(
                """
                SELECT record_date, machine_code,
                       COALESCE(SUM(planned_quantity), 0) AS planned_quantity,
                       COALESCE(SUM(actual_quantity), 0) AS actual_quantity
                FROM system_import.production_records
                WHERE record_date BETWEEN DATE_SUB(%s, INTERVAL 27 DAY) AND %s
                GROUP BY record_date, machine_code
                ORDER BY record_date, machine_code
                """,
                (selected_date, selected_date),
            )
            production_trends = cursor.fetchall()
            cursor.execute(
                """
                SELECT reports.event_date, reports.machine_code, reports.event_type,
                       COALESCE(options.display_name, reports.event_type) AS category_label,
                       COALESCE(SUM(reports.duration_minutes), 0) AS minutes
                FROM system_import.abnormality_reports AS reports
                LEFT JOIN system_import.event_type_options AS options
                  ON options.event_value = reports.event_type
                WHERE reports.event_date BETWEEN DATE_SUB(%s, INTERVAL 27 DAY) AND %s
                  AND LOWER(TRIM(reports.event_type)) NOT IN ('cleaning', 'normal production')
                GROUP BY reports.event_date, reports.machine_code, reports.event_type,
                         options.display_name
                ORDER BY reports.machine_code, reports.event_date, category_label
                """,
                (selected_date, selected_date),
            )
            abnormality_trends = cursor.fetchall()
    if uses_slitting_workflow():
        production = [slitting.enrich(row) for row in production]
        events = [slitting.enrich(row) for row in events]
        ongoing_events = [slitting.enrich(row) for row in ongoing_events]
    return {
        "report_date": record_date,
        "production": _json_safe_rows(production),
        "events": _json_safe_rows(events),
        "ongoing_events": _json_safe_rows(ongoing_events),
        "trend_dates": [
            (selected_date - timedelta(days=offset)).isoformat()
            for offset in range(27, -1, -1)
        ],
        "abnormality_trends": _json_safe_rows(abnormality_trends),
        "production_trends": _json_safe_rows(production_trends),
    }


def public_display_summary(record_date=None, shift_name=None):
    """Safe display summary: today task, last production day, cumulative machines."""
    try:
        slot_date, slot_shift = current_production_slot()
        selected_date = date.fromisoformat(record_date) if record_date else slot_date
        selected_shift = str(shift_name or slot_shift).title()
        allowed_shifts = (
            {"Day", "Night"} if uses_slitting_workflow() else {"Day", "Day 08-20", "Day 08-17"}
        )
        if selected_shift not in allowed_shifts:
            raise ValueError("Choose Day or Night shift.")
    except ValueError as exc:
        raise ValueError("Choose a valid display date.") from exc
    display_settings = get_display_settings()
    abnormality_chart_types = display_settings.get("abnormality_chart_types", [])
    type_placeholders = ", ".join(["%s"] * len(abnormality_chart_types))
    abnormality_type_clause = (
        f"AND reports.event_type IN ({type_placeholders})"
        if abnormality_chart_types
        else "AND 1 = 0"
    )
    with transaction() as cursor:
        cursor.execute(
            f"""
            SELECT record_date, shift_name, task_types, task_items, reminders, custom_task, formula_code, amount_needed, machine_assigned, updated_at
            FROM system_import.daily_tasks
            WHERE record_date = %s AND shift_name = %s
            """,
            (selected_date, selected_shift),
        )
        daily_task = cursor.fetchone()
        cursor.execute(
            """
            SELECT record_date, COUNT(*) AS record_count, COALESCE(SUM(planned_quantity), 0) AS planned,
                   COALESCE(SUM(actual_quantity), 0) AS actual,
                   COALESCE(SUM(qualified_quantity), 0) AS qualified,
                   COALESCE(SUM(qualified_pending), 0) AS qualified_pending,
                   MAX(updated_at) AS latest_production_update
            FROM system_import.production_records
            GROUP BY record_date
            ORDER BY record_date DESC
            LIMIT 1
            """,
        )
        production = cursor.fetchone()
        cursor.execute("""
            SELECT id, record_date, shift_name, formula_code, machine_code,
                   batch_number, actual_quantity
            FROM system_import.production_records
            WHERE qualified_pending = TRUE
            ORDER BY record_date, created_at, id
            """)
        pending_qualifications = cursor.fetchall()
        cursor.execute(
            """
            SELECT COUNT(*) AS total, COALESCE(SUM(state = 'open'), 0) AS open_events,
                   COALESCE(SUM(is_resolved = 'yes'), 0) AS resolved_events,
                   MAX(updated_at) AS latest_event_update
            FROM system_import.abnormality_reports WHERE event_date = %s
            """,
            (selected_date,),
        )
        events = cursor.fetchone()
        cursor.execute(
            """
            SELECT machine_code, COALESCE(SUM(planned_quantity), 0) AS planned,
                   COALESCE(SUM(actual_quantity), 0) AS actual,
                   COALESCE(SUM(qualified_pending), 0) AS pending
            FROM system_import.production_records
            GROUP BY machine_code ORDER BY machine_code
            """,
        )
        machines = cursor.fetchall()
        cursor.execute(
            """
            SELECT machine_code, COUNT(*) AS event_count, COALESCE(SUM(state = 'open'), 0) AS open_count
            FROM system_import.abnormality_reports
            GROUP BY machine_code
            """,
        )
        event_by_machine = {row["machine_code"]: row for row in cursor.fetchall()}
        cursor.execute(
            """
            SELECT event_date, start_time, end_time, shift_name, machine_code, machine_type,
                   event_type, severity, duration_minutes, is_resolved, state,
                   responsible_person, target_finish_date, effective_time_cost,
                   cost_failure_types, description, immediate_action
            FROM system_import.abnormality_reports
            WHERE event_date = %s
            ORDER BY event_date DESC, start_time DESC, created_at DESC
            LIMIT 12
            """,
            (selected_date,),
        )
        event_tracker = cursor.fetchall()
        cursor.execute("""
            SELECT id, event_date, start_time, shift_name, machine_code,
                   event_type, severity,
                   is_resolved, responsible_person, target_finish_date, custom_fields
            FROM system_import.abnormality_reports
            WHERE is_resolved = 'no'
            ORDER BY event_date DESC, start_time DESC, created_at DESC
            """)
        ongoing_events = cursor.fetchall()
        cursor.execute(
            """
            SELECT record_date, COALESCE(SUM(planned_quantity), 0) AS planned,
                   COALESCE(SUM(actual_quantity), 0) AS actual,
                   COALESCE(SUM(COALESCE(qualified_quantity, 0)), 0) AS qualified,
                   COALESCE(SUM(CASE WHEN qualified_pending = TRUE THEN actual_quantity ELSE 0 END), 0) AS pending
            FROM system_import.production_records
            WHERE record_date BETWEEN DATE_SUB(%s, INTERVAL 6 DAY) AND %s
            GROUP BY record_date ORDER BY record_date
            """,
            (selected_date, selected_date),
        )
        production_trend_rows = cursor.fetchall()
        cursor.execute(
            """
            SELECT record_date, machine_code, formula_code, planned_quantity,
                   actual_quantity, qualified_quantity, custom_fields
            FROM system_import.production_records
            WHERE record_date BETWEEN DATE_SUB(%s, INTERVAL 6 DAY) AND %s
            ORDER BY record_date, machine_code, formula_code
            """,
            (selected_date, selected_date),
        )
        weekly_production_detail_rows = cursor.fetchall()
        cursor.execute(
            f"""
            SELECT reports.event_date, reports.machine_code, reports.event_type,
                   COALESCE(options.display_name, reports.event_type) AS category_label,
                   COALESCE(SUM(reports.duration_minutes), 0) AS minutes
            FROM system_import.abnormality_reports AS reports
            LEFT JOIN system_import.event_type_options AS options
              ON options.event_value = reports.event_type
            WHERE reports.event_date BETWEEN DATE_SUB(%s, INTERVAL 6 DAY) AND %s
              {abnormality_type_clause}
            GROUP BY reports.event_date, reports.machine_code, reports.event_type, options.display_name
            ORDER BY reports.event_date, reports.machine_code, category_label
            """,
            (selected_date, selected_date, *abnormality_chart_types),
        )
        abnormality_trend_rows = cursor.fetchall()

    if daily_task:
        daily_task["task_items"] = json.loads(
            daily_task["task_items"] or "null"
        ) or _legacy_task_items(daily_task)
        daily_task["reminders"] = json.loads(daily_task["reminders"] or "[]")
        daily_task["shift_label"] = production_shift_label(selected_date, selected_shift)
        task_updated = daily_task["updated_at"]
    else:
        daily_task = {
            "record_date": selected_date.isoformat(),
            "shift_name": selected_shift,
            "shift_label": production_shift_label(selected_date, selected_shift),
            "task_items": [],
            "reminders": [],
            "updated_at": None,
        }
        task_updated = None
    if not production:
        production = {
            "record_date": None,
            "record_count": 0,
            "planned": 0,
            "actual": 0,
            "qualified": 0,
            "qualified_pending": 0,
            "latest_production_update": None,
        }

    for row in machines:
        event_row = event_by_machine.get(row["machine_code"], {})
        row["event_count"] = event_row.get("event_count", 0)
        row["open_count"] = event_row.get("open_count", 0)
    planned, actual, qualified = (
        float(production[key] or 0) for key in ("planned", "actual", "qualified")
    )
    latest = max(
        (
            value
            for value in (
                task_updated,
                production["latest_production_update"],
                events["latest_event_update"],
            )
            if value is not None
        ),
        default=None,
    )
    if uses_slitting_workflow():
        pending_qualifications = [
            r for r in slitting.summary()["records"] if r.get("qualified_quantity") is None
        ]
        for row in pending_qualifications:
            row["actual_quantity"] = row["input_rolls"] * row["stations"]
        ongoing_events = [slitting.enrich(row) for row in ongoing_events]
        weekly_production_details = [slitting.enrich(row) for row in weekly_production_detail_rows]
    else:
        weekly_production_details = []
    production_by_date = {row["record_date"]: row for row in production_trend_rows}
    abnormality_by_date = {}
    for row in abnormality_trend_rows:
        abnormality_by_date.setdefault(row["event_date"], []).append(
            {
                "type": row["event_type"],
                "label": row["category_label"],
                "machine": row["machine_code"],
                "minutes": float(row["minutes"] or 0),
            }
        )
    trend_dates = [selected_date - timedelta(days=offset) for offset in range(6, -1, -1)]
    production_trend = [
        {
            "date": day.isoformat(),
            "planned": float(production_by_date.get(day, {}).get("planned", 0) or 0),
            "actual": float(production_by_date.get(day, {}).get("actual", 0) or 0),
            "qualified": float(production_by_date.get(day, {}).get("qualified", 0) or 0),
            "pending": float(production_by_date.get(day, {}).get("pending", 0) or 0),
        }
        for day in trend_dates
    ]
    abnormality_trend = [
        {
            "date": day.isoformat(),
            "total_minutes": sum(
                segment["minutes"] for segment in abnormality_by_date.get(day, [])
            ),
            "categories": abnormality_by_date.get(day, []),
        }
        for day in trend_dates
    ]
    return {
        "display_date": selected_date.isoformat(),
        "daily_task": _json_safe_rows([daily_task])[0],
        "production": {
            "record_date": (
                production["record_date"].isoformat() if production["record_date"] else None
            ),
            "record_count": production["record_count"],
            "planned": planned,
            "actual": actual,
            "achievement_rate": round(actual / planned, 4) if planned else None,
            "qualified_rate": (
                round(qualified / actual, 4)
                if actual and not production["qualified_pending"]
                else None
            ),
            "qualified_pending": int(production["qualified_pending"] or 0),
        },
        "pending_qualifications": _json_safe_rows(pending_qualifications),
        "events": {
            "total": events["total"],
            "open": int(events["open_events"] or 0),
            "resolved": int(events["resolved_events"] or 0),
        },
        "machines": _json_safe_rows(machines),
        "event_tracker": _json_safe_rows(event_tracker),
        "ongoing_events": _json_safe_rows(ongoing_events),
        "trends": {
            "production": production_trend,
            "abnormalities": abnormality_trend,
            "production_details": _json_safe_rows(weekly_production_details),
        },
        "display_settings": display_settings,
        "last_updated": latest.isoformat(sep=" ", timespec="seconds") if latest else None,
    }


def manager_get_record(record_type, record_id):
    """Return one operational record for a manager-only correction screen."""
    table = {"production": "production_records", "abnormality": "abnormality_reports"}.get(
        record_type
    )
    if not table:
        raise ValueError("Unknown record type.")
    with transaction() as cursor:
        cursor.execute(f"SELECT * FROM system_import.{table} WHERE id = %s", (record_id,))
        row = cursor.fetchone()
    if not row:
        raise ValueError("Record not found.")
    if uses_slitting_workflow():
        row = slitting.enrich(row)
    return _json_safe_rows([row])[0]


def _manager_completion_values(existing, payload):
    """Validate completion only when this edit is actually completing the event."""
    requested_status = str(payload.get("is_resolved", existing["is_resolved"])).lower()
    if requested_status not in {"yes", "no"}:
        raise ValueError("Resolved status must be yes or no.")

    actual_finish_date = (
        payload.get("actual_finish_date") or existing.get("actual_finish_date") or None
    )
    solution_provided = str(
        payload.get("solution_provided") or existing.get("solution_provided") or ""
    ).strip()
    has_finish, has_solution = bool(actual_finish_date), bool(solution_provided)
    if has_finish != has_solution:
        raise ValueError("Actual finish date and solution provided must be completed together.")

    was_resolved = (
        str(existing.get("is_resolved", "no")).lower() == "yes"
        or existing.get("state") == "resolved"
    )
    if has_finish and has_solution:
        requested_status = "yes"
    elif requested_status == "yes" and not was_resolved:
        raise ValueError("Closing an event requires both actual finish date and solution provided.")

    return actual_finish_date, solution_provided, requested_status


def manager_update_record(
    record_type,
    record_id,
    payload,
    actor_id=None,
    audit_action="manager_corrected",
):
    """Correct a record while retaining its original date and audit history."""
    existing = manager_get_record(record_type, record_id)
    if uses_slitting_workflow():
        payload = (
            slitting.normalise_production(payload, existing)
            if record_type == "production"
            else slitting.normalise_abnormality(payload, existing)
        )
    with transaction() as cursor:
        if record_type == "production":
            planned = float(payload.get("planned_quantity", existing["planned_quantity"]))
            actual = float(payload.get("actual_quantity", existing["actual_quantity"]))
            raw_qualified = str(payload.get("qualified_quantity", "")).strip()
            qualified = float(raw_qualified) if raw_qualified else None
            if (
                planned < 0
                or actual < 0
                or (qualified is not None and (qualified < 0 or qualified > actual))
            ):
                raise ValueError("Production quantities are invalid.")
            cursor.execute(
                """UPDATE system_import.production_records SET shift_name=%s, machine_code=%s, formula_code=%s, batch_number=%s,
                planned_quantity=%s, actual_quantity=%s, qualified_quantity=%s, qualified_pending=%s, notes=%s,
                achievement_rate=%s, qualified_rate=%s, row_version=row_version+1 WHERE id=%s""",
                (
                    payload.get("shift_name", existing["shift_name"]),
                    payload.get("machine_code", existing["machine_code"]),
                    payload.get("formula_code", existing["formula_code"]),
                    payload.get("batch_number", existing["batch_number"]),
                    planned,
                    actual,
                    qualified,
                    qualified is None,
                    payload.get("notes", existing.get("notes") or ""),
                    actual / planned if planned else None,
                    qualified / actual if qualified is not None and actual else None,
                    record_id,
                ),
            )
        elif record_type == "abnormality":
            severity = str(payload.get("severity", existing["severity"])).lower()
            if severity not in {"normal", "low", "medium", "high"}:
                raise ValueError("Choose a valid severity or priority.")
            actual_finish_date, solution_provided, is_resolved = _manager_completion_values(
                existing, payload
            )
            if actual_finish_date and str(actual_finish_date) < str(existing["event_date"]):
                raise ValueError("Actual finish date cannot be earlier than the event date.")
            start_time = payload.get("start_time") or existing.get("start_time")
            end_time = payload.get("end_time") or existing.get("end_time")
            duration_minutes = _downtime_minutes(start_time, end_time)
            cursor.execute(
                """UPDATE system_import.abnormality_reports SET start_time=%s, end_time=%s, shift_name=%s, machine_code=%s,
                machine_type=%s, event_type=%s, severity=%s, duration_minutes=%s, description=%s, immediate_action=%s,
                responsible_person=%s, target_finish_date=%s, actual_finish_date=%s, solution_provided=%s,
                effective_time_cost=%s, is_resolved=%s, state=%s, row_version=row_version+1 WHERE id=%s""",
                (
                    start_time,
                    end_time,
                    payload.get("shift_name", existing["shift_name"]),
                    payload.get("machine_code", existing["machine_code"]),
                    payload.get("machine_type", existing["machine_type"]),
                    payload.get("event_type", existing["event_type"]),
                    severity,
                    duration_minutes,
                    payload.get("description", existing.get("description") or ""),
                    payload.get("immediate_action", existing.get("immediate_action") or ""),
                    payload.get("responsible_person", existing.get("responsible_person") or ""),
                    payload.get("target_finish_date") or None,
                    actual_finish_date,
                    solution_provided,
                    payload.get("effective_time_cost") or None,
                    is_resolved,
                    "resolved" if is_resolved == "yes" else "open",
                    record_id,
                ),
            )
        else:
            raise ValueError("Unknown record type.")
        if uses_slitting_workflow():
            table = "production_records" if record_type == "production" else "abnormality_reports"
            cursor.execute(
                f"UPDATE system_import.{table} SET custom_fields=%s WHERE id=%s",
                (json.dumps(payload["custom_fields"]), record_id),
            )
        _audit(
            cursor,
            actor_id,
            f"{record_type}_record",
            record_id,
            audit_action,
            payload,
        )
    return manager_get_record(record_type, record_id)
