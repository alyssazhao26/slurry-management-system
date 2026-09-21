from contextlib import contextmanager
from datetime import date, datetime

import pytest
from flask import Flask

from app.services import attendance


def test_employee_pin_must_be_unique(monkeypatch):
    class Cursor:
        def execute(self, sql, params=None):
            self.sql = sql

        def fetchone(self):
            return {"id": 7}

    @contextmanager
    def fake_transaction():
        yield Cursor()

    monkeypatch.setattr(attendance, "transaction", fake_transaction)
    monkeypatch.setattr(attendance.bcrypt, "hashpw", lambda value, salt: b"new-hash")
    monkeypatch.setattr(attendance.bcrypt, "gensalt", lambda **kwargs: b"salt")

    with Flask(__name__).app_context(), pytest.raises(ValueError, match="already assigned"):
        attendance.save_employee({
            "staffing_company": "Agency",
            "employee_name": "Worker",
            "pin": "1234",
        })


@pytest.mark.parametrize("pin", ["123", "12345", "12a4"])
def test_employee_pin_must_be_exactly_four_digits(pin):
    with pytest.raises(ValueError, match="exactly 4 digits"):
        attendance.save_employee({"staffing_company": "Agency", "employee_name": "Worker", "pin": pin})


def test_manager_employee_list_includes_visible_pin_column():
    source = __import__("pathlib").Path(attendance.__file__).read_text(encoding="utf-8")
    assert "job_title,pin_code,pin_reset_required,is_active" in source


def test_manager_employee_action_hides_reset_label_after_setup():
    root = __import__("pathlib").Path(attendance.__file__).parents[1]
    source = (root / "static" / "app.js").read_text(encoding="utf-8")
    assert "Boolean(Number(row.pin_reset_required))" in source
    assert "Edit employee / 编辑员工" in source
    assert "Edit / Reset PIN / 编辑或重置密码" not in source


@pytest.mark.parametrize(
    "database_value, expected",
    [(0, False), (1, True), (False, False), (True, True)],
)
def test_employee_reset_state_is_normalised_to_boolean(database_value, expected):
    row = {
        "pin_reset_required": database_value,
        "is_active": 1,
    }
    result = attendance._normalise_employee_row(row)
    assert result["pin_reset_required"] is expected
    assert result["is_active"] is True


def test_first_login_pin_reset_updates_visible_pin(monkeypatch):
    executed = []
    class Cursor:
        fetch_count = 0
        def execute(self, sql, params=None): executed.append((sql, params))
        def fetchone(self):
            self.fetch_count += 1
            if self.fetch_count == 1:
                return {"id": 4, "pin_hash": "temporary-hash", "pin_reset_required": True}
            return None
    @contextmanager
    def fake_transaction(): yield Cursor()
    monkeypatch.setattr(attendance, "transaction", fake_transaction)
    monkeypatch.setattr(attendance.bcrypt, "hashpw", lambda value, salt: b"personal-hash")
    monkeypatch.setattr(attendance.bcrypt, "gensalt", lambda **kwargs: b"salt")
    monkeypatch.setattr(attendance.bcrypt, "checkpw", lambda value, hashed: False)
    with Flask(__name__).app_context():
        attendance.reset_employee_pin(4, "4826", "4826")
    update = next(params for sql, params in executed if "pin_reset_required=FALSE" in sql)
    assert update == ("personal-hash", "4826", 4)


def test_first_login_pin_must_differ_from_temporary_pin(monkeypatch):
    class Cursor:
        def execute(self, sql, params=None): pass
        def fetchone(self): return {"id": 4, "pin_hash": "temporary-hash", "pin_reset_required": True}
    @contextmanager
    def fake_transaction(): yield Cursor()
    monkeypatch.setattr(attendance, "transaction", fake_transaction)
    monkeypatch.setattr(attendance.bcrypt, "hashpw", lambda value, salt: b"personal-hash")
    monkeypatch.setattr(attendance.bcrypt, "gensalt", lambda **kwargs: b"salt")
    monkeypatch.setattr(attendance.bcrypt, "checkpw", lambda value, hashed: True)
    with Flask(__name__).app_context(), pytest.raises(ValueError, match="different from the temporary"):
        attendance.reset_employee_pin(4, "4826", "4826")


def test_attendance_login_timestamp_is_serialised_for_display():
    rows = attendance._serialise_attendance_rows([{
        "attendance_date": date(2026, 9, 14),
        "login_timestamp": datetime(2026, 9, 14, 7, 30),
    }])
    assert rows[0] == {
        "attendance_date": "2026-09-14",
        "login_timestamp": "2026-09-14T07:30",
    }


@pytest.mark.parametrize("field,value", [
    ("paid_break_hours", "0.25"),
    ("paid_break_hours", "2.5"),
    ("total_hours", "2.5"),
    ("total_hours", "12.5"),
    ("overtime_hours", "0.25"),
])
def test_attendance_hours_require_valid_half_hour_values(field, value):
    payload = {"row_version": 0, field: value}
    with pytest.raises(ValueError, match="0.5-hour increments"):
        attendance.update_attendance(1, payload)


def test_attendance_filters_include_partial_employee_name():
    where, params = attendance._filters(
        staffing_company="Onin",
        employee_name="Test",
    )
    assert "e.staffing_company LIKE %s" in where
    assert "e.employee_name LIKE %s" in where
    assert params == ["%Onin%", "%Test%"]


def test_attendance_login_passes_only_pin(monkeypatch):
    from app import create_app
    from app.config import Config

    monkeypatch.setattr("app.db.verify_ui_schema", lambda: None)
    monkeypatch.setattr(Config, "ENVIRONMENT", "development")
    monkeypatch.setattr(Config, "BEHIND_TLS_PROXY", False)
    monkeypatch.setattr(Config, "SESSION_COOKIE_SECURE", False)
    monkeypatch.setattr(Config, "ALLOWED_HOSTS", {"localhost"})
    received = []
    monkeypatch.setattr("app.routes.authenticate_employee", lambda pin: received.append(pin) or {"id": 4, "employee_name": "Worker"})
    monkeypatch.setattr("app.routes.record_employee_login", lambda employee_id: None)
    monkeypatch.setattr("app.routes.employee_history", lambda employee_id: [])
    client = create_app().test_client()
    token = client.get("/api/csrf").json["token"]
    response = client.post("/api/attendance/login", json={"pin": "4826"}, headers={"X-CSRF-Token": token})
    assert response.status_code == 200
    assert received == ["4826"]


@pytest.mark.parametrize("pin", ["", "123", "12345", "abcd"])
def test_attendance_authentication_rejects_non_four_digit_pin(pin):
    with pytest.raises(ValueError, match="exactly 4 digits"):
        attendance.authenticate_employee(pin)


def test_attendance_login_rate_limits_failed_pins(monkeypatch):
    from app import create_app
    from app.config import Config

    monkeypatch.setattr("app.db.verify_ui_schema", lambda: None)
    monkeypatch.setattr(Config, "ENVIRONMENT", "development")
    monkeypatch.setattr(Config, "BEHIND_TLS_PROXY", False)
    monkeypatch.setattr(Config, "SESSION_COOKIE_SECURE", False)
    monkeypatch.setattr(Config, "ALLOWED_HOSTS", {"localhost"})
    monkeypatch.setattr(Config, "MANAGER_LOGIN_MAX_ATTEMPTS", 2)
    monkeypatch.setattr("app.routes.authenticate_employee", lambda pin: (_ for _ in ()).throw(ValueError("Invalid PIN.")))
    client = create_app().test_client()
    token = client.get("/api/csrf").json["token"]
    request = {"json": {"pin": "0000"}, "headers": {"X-CSRF-Token": token}}
    assert client.post("/api/attendance/login", **request).status_code == 401
    assert client.post("/api/attendance/login", **request).status_code == 401
    assert client.post("/api/attendance/login", **request).status_code == 429


def test_employee_submits_daily_hour_values(monkeypatch):
    executed = []
    class Cursor:
        def execute(self, sql, params=None): executed.append((sql, params))
        def fetchone(self): return {"id": 9}
    @contextmanager
    def fake_transaction(): yield Cursor()
    monkeypatch.setattr(attendance, "transaction", fake_transaction)
    with Flask(__name__).app_context():
        attendance.submit_attendance(4, {"paid_break_hours": "0.5", "total_hours": "8.0", "overtime_hours": "1.5"})
    update = next((params for sql, params in executed if "UPDATE system_import.attendance_records" in sql), None)
    assert update[:3] == (0.5, 8.0, 1.5)


def test_employee_session_expires_after_five_minutes(monkeypatch):
    from app import create_app
    from app.config import Config
    clock = {"now": 1000}
    monkeypatch.setattr("app.db.verify_ui_schema", lambda: None)
    monkeypatch.setattr(Config, "ENVIRONMENT", "development")
    monkeypatch.setattr(Config, "BEHIND_TLS_PROXY", False)
    monkeypatch.setattr(Config, "SESSION_COOKIE_SECURE", False)
    monkeypatch.setattr(Config, "ALLOWED_HOSTS", {"localhost"})
    monkeypatch.setattr("app.routes.authenticate_employee", lambda pin: {"id": 4, "employee_name": "Worker"})
    monkeypatch.setattr("app.routes.record_employee_login", lambda employee_id: None)
    monkeypatch.setattr("app.routes.employee_history", lambda employee_id: [])
    monkeypatch.setattr("app.routes.time", lambda: clock["now"])
    client = create_app().test_client()
    token = client.get("/api/csrf").json["token"]
    assert client.post("/api/attendance/login", json={"pin": "4826"}, headers={"X-CSRF-Token": token}).status_code == 200
    clock["now"] = 1301
    assert client.get("/api/attendance/me").status_code == 401


def test_opening_attendance_page_requires_new_login(monkeypatch):
    from app import create_app
    from app.config import Config

    monkeypatch.setattr("app.db.verify_ui_schema", lambda: None)
    monkeypatch.setattr(Config, "ENVIRONMENT", "development")
    monkeypatch.setattr(Config, "BEHIND_TLS_PROXY", False)
    monkeypatch.setattr(Config, "SESSION_COOKIE_SECURE", False)
    monkeypatch.setattr(Config, "ALLOWED_HOSTS", {"localhost"})
    monkeypatch.setattr(
        "app.routes.authenticate_employee",
        lambda pin: {"id": 4, "employee_name": "Worker"},
    )
    monkeypatch.setattr("app.routes.record_employee_login", lambda employee_id: None)
    monkeypatch.setattr("app.routes.employee_history", lambda employee_id: [])
    client = create_app().test_client()
    token = client.get("/api/csrf").json["token"]
    response = client.post(
        "/api/attendance/login",
        json={"pin": "4826"},
        headers={"X-CSRF-Token": token},
    )
    assert response.status_code == 200
    assert client.get("/api/attendance/me").status_code == 200

    assert client.get("/attendance").status_code == 200
    assert client.get("/api/attendance/me").status_code == 401


def test_employee_page_replaces_clock_buttons_with_hour_selectors():
    root = __import__("pathlib").Path(attendance.__file__).parents[1]
    template = (root / "templates" / "attendance.html").read_text(encoding="utf-8")
    assert "paid_break_hours" in template and "total_hours" in template and "overtime_hours" in template
    assert "Clock in / 上班打卡" not in template
    assert "5 minutes after login" in template
    assert "Every new page visit requires login" in template
    assert "first-login-pin-form" in template
    assert "Set your personal PIN" in template


def test_manager_attendance_uses_one_login_timestamp_and_add_record():
    root = __import__("pathlib").Path(attendance.__file__).parents[1]
    source = (root / "static" / "app.js").read_text(encoding="utf-8")
    dashboard = (root / "templates" / "dashboard.html").read_text(encoding="utf-8")

    assert "Login timestamp / 登录时间" in source
    assert 'name="clock_in"' not in source
    assert 'name="clock_out"' not in source
    assert "new-attendance-record" in dashboard
    assert "attendance-record-form" in dashboard
    record_dialog = dashboard.split('id="attendance-record-dialog"', 1)[1]
    record_dialog = record_dialog.split('</dialog>', 1)[0]
    assert "Status / 状态" not in record_dialog
    assert "attendance-employee-options" in source


def test_employee_login_creates_first_daily_timestamp(monkeypatch):
    executed = []

    class Cursor:
        def execute(self, sql, params=None):
            executed.append((sql, params))

    @contextmanager
    def fake_transaction():
        yield Cursor()

    monkeypatch.setattr(attendance, "transaction", fake_transaction)
    monkeypatch.setattr(attendance, "datetime", type(
        "FixedDateTime",
        (),
        {"now": staticmethod(lambda: datetime(2026, 9, 15, 8, 12))},
    ))
    with Flask(__name__).app_context():
        attendance.record_employee_login(4)

    sql, params = executed[0]
    assert "login_timestamp=COALESCE" in sql
    assert params == (4, date(2026, 9, 15), datetime(2026, 9, 15, 8, 12))


def test_attendance_hidden_state_cannot_be_overridden_by_form_layout():
    root = __import__("pathlib").Path(attendance.__file__).parents[1]
    stylesheet = (root / "static" / "attendance.css").read_text(encoding="utf-8")
    assert "[hidden] { display: none !important; }" in stylesheet


def test_attendance_export_is_employee_by_date_matrix(monkeypatch):
    from openpyxl import load_workbook

    rows = [
        {
            "staffing_company": "GNEM",
            "employee_name": "Britney Anderson",
            "department": "Slurry",
            "attendance_date": "2026-09-14",
            "total_hours": 8,
            "paid_break_hours": 0.5,
            "overtime_hours": 1,
        },
        {
            "staffing_company": "GNEM",
            "employee_name": "Britney Anderson",
            "department": "Slurry",
            "attendance_date": "2026-09-15",
            "total_hours": 8,
            "paid_break_hours": 1,
            "overtime_hours": None,
        },
    ]
    monkeypatch.setattr(attendance, "list_attendance", lambda **kwargs: (rows, 2))

    workbook = load_workbook(attendance.export_attendance())
    assert workbook.sheetnames == ["General", "Slurry"]
    sheet = workbook["General"]

    assert sheet["A1"].value == "Staffing company / 劳务公司"
    assert sheet["B1"].value == "Employee name / 员工姓名"
    assert sheet["C1"].value.date() == date(2026, 9, 14)
    assert sheet["C1"].number_format == "d"
    assert sheet["C2"].value == "Mon"
    assert sheet["A3"].value == "GNEM"
    assert sheet["B3"].value == "Britney Anderson"
    assert sheet["C3"].value == 8.5
    assert sheet["D3"].value == 7
    assert workbook["Slurry"]["B3"].value == "Britney Anderson"


def test_attendance_export_keeps_weekend_days_separate_and_uses_zero(monkeypatch):
    from openpyxl import load_workbook

    rows = [
        {
            "staffing_company": "GNEM",
            "employee_name": "Employee A",
            "department": "Slurry",
            "attendance_date": "2026-09-18",
            "total_hours": 8,
            "paid_break_hours": 0,
            "overtime_hours": 0,
        },
        {
            "staffing_company": "GNEM",
            "employee_name": "Employee A",
            "department": "Slurry",
            "attendance_date": "2026-09-21",
            "total_hours": 8,
            "paid_break_hours": 0,
            "overtime_hours": 0,
        },
    ]
    monkeypatch.setattr(attendance, "list_attendance", lambda **kwargs: (rows, 2))

    sheet = load_workbook(attendance.export_attendance())["General"]
    assert [sheet.cell(2, column).value for column in range(3, 7)] == ["Fri", "Sat", "Sun", "Mon"]
    assert sheet["D3"].value == 0
    assert sheet["E3"].value == 0


def test_attendance_export_sorts_employee_names_alphabetically(monkeypatch):
    from openpyxl import load_workbook

    rows = [
        {
            "staffing_company": "Agency B",
            "employee_name": "Zoey Chen",
            "department": "Slurry",
            "attendance_date": "2026-09-21",
            "total_hours": 8,
            "paid_break_hours": 0,
            "overtime_hours": 0,
        },
        {
            "staffing_company": "Agency A",
            "employee_name": "amy Li",
            "department": "Slurry",
            "attendance_date": "2026-09-21",
            "total_hours": 8,
            "paid_break_hours": 0,
            "overtime_hours": 0,
        },
    ]
    monkeypatch.setattr(attendance, "list_attendance", lambda **kwargs: (rows, 2))

    workbook = load_workbook(attendance.export_attendance())
    assert [workbook["General"][cell].value for cell in ("B3", "B4")] == ["amy Li", "Zoey Chen"]
    assert [workbook["Slurry"][cell].value for cell in ("B3", "B4")] == ["amy Li", "Zoey Chen"]


def test_attendance_export_creates_safe_unique_department_sheets(monkeypatch):
    from openpyxl import load_workbook

    rows = [
        {
            "staffing_company": "Agency",
            "employee_name": "Employee A",
            "department": "Slurry/Production",
            "attendance_date": "2026-09-15",
            "total_hours": 8,
            "paid_break_hours": 0.5,
            "overtime_hours": 0,
        },
        {
            "staffing_company": "Agency",
            "employee_name": "Employee B",
            "department": "Slurry:Production",
            "attendance_date": "2026-09-15",
            "total_hours": 9,
            "paid_break_hours": 1,
            "overtime_hours": 0.5,
        },
    ]
    monkeypatch.setattr(attendance, "list_attendance", lambda **kwargs: (rows, 2))

    workbook = load_workbook(attendance.export_attendance())

    assert workbook.sheetnames == [
        "General",
        "Slurry-Production",
        "Slurry-Production (2)",
    ]
    assert workbook["Slurry-Production"]["B3"].value == "Employee A"
    assert workbook["Slurry-Production (2)"]["B3"].value == "Employee B"
